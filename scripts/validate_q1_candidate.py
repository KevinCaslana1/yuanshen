"""Validate a generated Q1 result1.xlsx candidate against the frozen contract.

This module is intentionally validation-only. It never creates or modifies a
workbook, and it fails closed when the candidate is absent.
"""

from __future__ import annotations

import argparse
import json
import math
import sys
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

try:
    sys.stdout.reconfigure(encoding="utf-8")
except AttributeError:
    pass

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - environment failure
    raise SystemExit("openpyxl is required for Q1 candidate validation") from exc


PRECISION = 4
EXPECTED_ROWS = 1801
EXPECTED_COLUMNS = 22
EXPECTED_TIMES = tuple(range(1, 1801))
EXPECTED_DISTANCES = tuple(Decimal(i) / Decimal(10) for i in range(21))


def _decimal(value: Any) -> Decimal | None:
    """Convert a cell value to Decimal, rejecting booleans and non-finite values."""

    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        return None
    if isinstance(value, float) and not math.isfinite(value):
        return None
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None
    return result if result.is_finite() else None


def _same_decimal(actual: Any, expected: Decimal) -> bool:
    value = _decimal(actual)
    return value is not None and value == expected


def _reported_value_ok(value: Any) -> bool:
    """Require a finite numeric cell whose value is already 4-decimal HALF_UP output."""

    decimal_value = _decimal(value)
    if decimal_value is None:
        return False
    rounded = decimal_value.quantize(Decimal("1e-4"), rounding=ROUND_HALF_UP)
    return decimal_value == rounded


def _q1_paths(root: Path) -> tuple[Path, Path]:
    config_path = root / "config" / "deliverables.json"
    config = json.loads(config_path.read_text(encoding="utf-8"))
    q1 = next(item for item in config["deliverables"] if item["question"] == "Q1")
    return root / q1["candidate_path"], root / q1["official_template"]


def validate_candidate(root: Path, candidate_path: Path | None = None) -> list[str]:
    """Return all contract/data errors for a Q1 candidate workbook."""

    errors: list[str] = []
    configured_candidate, official_template = _q1_paths(root)
    candidate = (candidate_path or configured_candidate).resolve()
    candidate_root = (root / "deliverables" / "candidate").resolve()
    official_root = (root / "A题").resolve()

    try:
        candidate.relative_to(candidate_root)
    except ValueError:
        errors.append(f"candidate path escapes deliverables/candidate: {candidate}")
    try:
        candidate.relative_to(official_root)
        errors.append("candidate path overlaps A题 official source")
    except ValueError:
        pass

    if not candidate.is_file():
        return errors + [f"missing Q1 candidate workbook: {candidate}"]

    try:
        workbook = openpyxl.load_workbook(candidate, read_only=True, data_only=False)
    except Exception as exc:
        return errors + [f"cannot open Q1 candidate workbook: {exc}"]

    try:
        expected_sheets = ("温度", "水分浓度")
        actual_sheets = tuple(workbook.sheetnames)
        if actual_sheets != expected_sheets:
            errors.append(f"sheet mismatch: expected={expected_sheets} actual={actual_sheets}")

        try:
            template = openpyxl.load_workbook(official_template, read_only=True, data_only=False)
        except Exception as exc:
            errors.append(f"cannot open official Q1 template: {exc}")
            template = None

        for sheet_name in expected_sheets:
            if sheet_name not in workbook.sheetnames:
                errors.append(f"missing sheet: {sheet_name}")
                continue
            sheet = workbook[sheet_name]
            if sheet.max_row != EXPECTED_ROWS or sheet.max_column != EXPECTED_COLUMNS:
                errors.append(
                    f"{sheet_name} shape mismatch: expected={EXPECTED_ROWS}x{EXPECTED_COLUMNS} "
                    f"actual={sheet.max_row}x{sheet.max_column}"
                )

            if template is not None and sheet_name in template.sheetnames:
                template_sheet = template[sheet_name]
                if sheet["A1"].value != template_sheet["A1"].value:
                    errors.append(f"{sheet_name}!A1 does not match official template header")

            for column, distance in enumerate(EXPECTED_DISTANCES, start=2):
                if not _same_decimal(sheet.cell(1, column).value, distance):
                    errors.append(
                        f"{sheet_name}!{sheet.cell(1, column).coordinate} distance header "
                        f"is not exactly {distance} cm"
                    )

            for row, expected_time in enumerate(EXPECTED_TIMES, start=2):
                if not _same_decimal(sheet.cell(row, 1).value, Decimal(expected_time)):
                    errors.append(
                        f"{sheet_name}!A{row} time is not the strict 1..1800 sequence"
                    )

            for row in range(2, EXPECTED_ROWS + 1):
                for column in range(2, EXPECTED_COLUMNS + 1):
                    value = sheet.cell(row, column).value
                    if not _reported_value_ok(value):
                        errors.append(
                            f"{sheet_name}!{sheet.cell(row, column).coordinate} is not a finite "
                            "numeric value already rounded to 4 decimals (ROUND_HALF_UP)"
                        )
                        if len(errors) >= 100:
                            errors.append("error list truncated after 100 issues")
                            return errors

        if template is not None:
            template.close()
    finally:
        workbook.close()

    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--candidate", type=Path, default=None)
    args = parser.parse_args()
    errors = validate_candidate(args.root.resolve(), args.candidate)
    if errors:
        for error in errors:
            print(f"[FAIL] {error}")
        print(f"Q1 candidate validation failed: {len(errors)} issue(s)")
        return 1
    print("[PASS] Q1 candidate matches the frozen 1801x22 result1 contract")
    print("[PASS] axes, sheet names, finite numeric cells, and 4-decimal output precision validated")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
