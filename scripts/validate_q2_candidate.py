"""Validate the Q2 result2 candidate against the frozen production contract."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import random
import sys
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

try:
    sys.stdout.reconfigure(encoding="utf-8")
except AttributeError:
    pass

ROOT = Path(__file__).resolve().parents[1]
FREEZE = ROOT / "experiments" / "Q2_FREEZE_RUN"
MANIFEST = ROOT / "experiments" / "Q2_CANONICAL_DATA_MANIFEST.json"
CANDIDATE = ROOT / "deliverables" / "candidate" / "result2.xlsx"
TEMPLATE = ROOT / "A题" / "附件" / "附件3" / "result2.xlsx"
EXPECTED_SHEETS = (("温度", "temperature_C"), ("水分浓度", "moisture_kg_kg"))
PAPER_TIMES = (1800, 3600, 5400, 7200, 9000, 10800)
PAPER_RADII = (0.0, 0.5, 1.0, 1.5, 2.0)
RANDOM_SEED = 20260912

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required for Q2 candidate validation") from exc

from src.q2.lineage import production_csv_path


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def rounded(value: float) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)


def decimal_value(value: Any) -> Decimal | None:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        return None
    if isinstance(value, float) and not math.isfinite(value):
        return None
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None
    return result if result.is_finite() else None


def validate_source_contract(config: dict[str, Any], hashes: dict[str, Any]) -> Path:
    source = production_csv_path("experiments/Q2_FREEZE_RUN/run_1/official_samples.csv")
    expected = hashes.get("run_1", {}).get("sampled_output")
    if expected != sha256(source):
        raise ValueError("run_1 sampled source hash does not match output_hashes.json")
    if config.get("official_sample_rule", "").find("time_s=1..final_horizon") < 0:
        raise ValueError("frozen official sample rule is missing t=1..final_horizon")
    return source


def collect_source_trace(source: Path, wanted: set[tuple[int, float, str]], expected_horizon: int) -> dict[tuple[int, float, str], float]:
    result: dict[tuple[int, float, str], float] = {}
    expected_rows = expected_horizon * 21
    with source.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        if list(reader.fieldnames or ()) != ["time_s", "radius_cm", "temperature_K", "temperature_C", "moisture_kg_kg"]:
            raise ValueError("run_1 sampled source header mismatch")
        rows = 0
        for row in reader:
            time_s = int(round(float(row["time_s"])))
            radius = round(float(row["radius_cm"]), 1)
            for field in ("temperature_C", "moisture_kg_kg"):
                key = (time_s, radius, field)
                if key in wanted:
                    result[key] = float(row[field])
            rows += 1
        if rows != expected_rows:
            raise ValueError(f"run_1 sampled source row count mismatch: {rows} != {expected_rows}")
    missing = wanted - result.keys()
    if missing:
        raise ValueError(f"source trace points missing: {sorted(missing)[:5]}")
    return result


def validate_workbook(path: Path, horizon: int, source_trace: dict[tuple[int, float, str], float], random_indices: list[int]) -> dict[str, Any]:
    errors: list[str] = []
    if path.name != "result2.xlsx":
        errors.append("candidate filename is not result2.xlsx")
    try:
        path.resolve().relative_to((ROOT / "deliverables" / "candidate").resolve())
    except ValueError:
        errors.append("candidate is outside deliverables/candidate")
    if path.resolve().is_relative_to((ROOT / "A题").resolve()):
        errors.append("candidate overlaps immutable A题 official source")
    if not path.is_file():
        return {"status": "FAIL", "errors": errors + [f"missing candidate: {path}"]}

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if tuple(workbook.sheetnames) != tuple(name for name, _ in EXPECTED_SHEETS):
            errors.append(f"sheet mismatch: {workbook.sheetnames}")
        template = openpyxl.load_workbook(TEMPLATE, read_only=True, data_only=False)
        try:
            random_targets = set(random_indices)
            trace: list[dict[str, Any]] = []
            paper_trace: dict[str, list[dict[str, Any]]] = {field: [] for _, field in EXPECTED_SHEETS}
            for sheet_name, field in EXPECTED_SHEETS:
                if sheet_name not in workbook.sheetnames:
                    errors.append(f"missing sheet: {sheet_name}")
                    continue
                sheet = workbook[sheet_name]
                if sheet.max_row != horizon + 1 or sheet.max_column != 22:
                    errors.append(f"{sheet_name} shape={sheet.max_row}x{sheet.max_column}, expected={(horizon + 1)}x22")
                if sheet["A1"].value != template[sheet_name]["A1"].value:
                    errors.append(f"{sheet_name}!A1 does not match official template")
                rows = sheet.iter_rows(min_row=1, max_row=horizon + 1, min_col=1, max_col=22, values_only=True)
                header = tuple(next(rows, ()))
                for col, radius in enumerate((Decimal(i) / Decimal(10) for i in range(21)), start=1):
                    actual = decimal_value(header[col]) if len(header) > col else None
                    if actual != radius:
                        errors.append(f"{sheet_name} header column {col + 1} is {actual}, expected {radius}")
                for row_index, values in enumerate(rows, start=2):
                    time_value = decimal_value(values[0] if values else None)
                    expected_time = Decimal(row_index - 1)
                    if time_value != expected_time:
                        errors.append(f"{sheet_name}!A{row_index} time={time_value}, expected={expected_time}")
                    if len(values) < 22:
                        errors.append(f"{sheet_name} row {row_index} has fewer than 22 columns")
                        continue
                    time_int = row_index - 1
                    for pos_index, value in enumerate(values[1:22]):
                        numeric = decimal_value(value)
                        if numeric is None:
                            errors.append(f"{sheet_name} row {row_index} col {pos_index + 2} is blank/non-finite")
                            continue
                        if numeric != numeric.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP):
                            errors.append(f"{sheet_name} row {row_index} col {pos_index + 2} is not 4-decimal ROUND_HALF_UP")
                        sample_index = (time_int - 1) * 21 + pos_index
                        if sample_index in random_targets:
                            radius = round(pos_index * 0.1, 1)
                            raw = source_trace[(time_int, radius, field)]
                            expected = rounded(raw)
                            trace.append({"sheet": sheet_name, "time_s": time_int, "radius_cm": radius, "candidate": str(numeric), "source_raw": repr(raw), "expected_round_half_up": str(expected), "pass": numeric == expected})
                            if numeric != expected:
                                errors.append(f"random trace mismatch {sheet_name} t={time_int} r={radius}")
                        if time_int in PAPER_TIMES and round(pos_index * 0.1, 1) in PAPER_RADII:
                            radius = round(pos_index * 0.1, 1)
                            raw = source_trace[(time_int, radius, field)]
                            expected = rounded(raw)
                            item = {"time_s": time_int, "radius_cm": radius, "candidate": str(numeric), "source_raw": repr(raw), "expected_round_half_up": str(expected), "pass": numeric == expected}
                            paper_trace[field].append(item)
                            if numeric != expected:
                                errors.append(f"paper trace mismatch {sheet_name} t={time_int} r={radius}")
                    if len(errors) >= 100:
                        errors.append("error list truncated after 100 issues")
                        break
                if len(errors) >= 100:
                    break
        finally:
            template.close()
    finally:
        workbook.close()
    random_pass = sum(item["pass"] for item in trace)
    paper_summary = {field: {"count": len(items), "pass_count": sum(item["pass"] for item in items)} for field, items in paper_trace.items()}
    return {"status": "PASS" if not errors and len(trace) == len(random_indices) * 2 and all(item["pass_count"] == 30 for item in paper_summary.values()) else "FAIL", "errors": errors, "random_seed": RANDOM_SEED, "random_cell_count_per_sheet": len(random_indices), "random_trace_pass_count": random_pass, "random_trace_total": len(random_indices) * 2, "paper_tables": paper_summary, "trace": trace}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--candidate", type=Path, default=CANDIDATE)
    args = parser.parse_args()
    config = json.loads((FREEZE / "config.json").read_text(encoding="utf-8"))
    hashes = json.loads((FREEZE / "output_hashes.json").read_text(encoding="utf-8"))
    horizon = int(config["final_horizon_s"])
    source = validate_source_contract(config, hashes)
    rng = random.Random(RANDOM_SEED)
    random_indices = rng.sample(range(horizon * 21), 30)
    wanted = {(time_s, round(index * 0.1, 1), field) for field in ("temperature_C", "moisture_kg_kg") for time_s, index in [(sample_index // 21 + 1, sample_index % 21) for sample_index in random_indices]}
    wanted.update((time_s, radius, field) for field in ("temperature_C", "moisture_kg_kg") for time_s in PAPER_TIMES for radius in PAPER_RADII)
    source_trace = collect_source_trace(source, wanted, horizon)
    validation = validate_workbook(args.candidate.resolve(), horizon, source_trace, random_indices)
    validation["candidate_sha256"] = sha256(args.candidate.resolve()) if args.candidate.is_file() else None
    validation["source"] = {"path": "experiments/Q2_FREEZE_RUN/run_1/official_samples.csv", "sha256": sha256(source), "lineage_manifest": str(MANIFEST.relative_to(ROOT)).replace("\\", "/")}
    (FREEZE / "candidate_validation.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({key: validation[key] for key in ("status", "candidate_sha256", "random_trace_pass_count", "random_trace_total", "paper_tables", "errors")}, ensure_ascii=False, indent=2))
    return 0 if validation["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
