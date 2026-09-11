"""Generate and audit the Q1 candidate workbook from Q1_FREEZE_RUN only."""

from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import pickle
import random
import shutil
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Dict, Iterable, Tuple

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.validate_inputs import sha256
from scripts.validate_q1_candidate import validate_candidate
from src.common.paths import assert_not_official_output

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required for Q1 candidate generation") from exc


FREEZE_DIR = ROOT / "experiments" / "Q1_FREEZE_RUN"
SOURCE_REFERENCE = FREEZE_DIR / "run_1" / "internal_output_reference.pkl.gz"
OFFICIAL_TEMPLATE = ROOT / "A题" / "附件" / "附件3" / "result1.xlsx"
CANDIDATE = ROOT / "deliverables" / "candidate" / "result1.xlsx"
OUTPUT_POSITIONS_M = tuple(index * 0.001 for index in range(21))
PAPER_TIMES_S = (100, 300, 600, 900, 1200, 1500, 1800)
PAPER_POSITIONS_M = (0.0, 0.005, 0.01, 0.015, 0.02)
EXPECTED_SHEETS = (("温度", "temperature_c"), ("水分浓度", "moisture_kg_kg"))
QUANTUM = Decimal("0.0001")


def _load_payload(path: Path) -> Dict[str, Any]:
    with gzip.open(path, "rb") as stream:
        payload = pickle.load(stream)
    if payload.get("schema_version") != "Q1_FREEZE_INTERNAL_OUTPUT_V1":
        raise ValueError("unsupported or untrusted freeze reference schema")
    return payload


def _rounded(value: float) -> float:
    return float(Decimal(str(value)).quantize(QUANTUM, rounding=ROUND_HALF_UP))


def _position_index(nodes: Iterable[float], position_m: float) -> int:
    nodes = tuple(nodes)
    index = min(range(len(nodes)), key=lambda candidate: abs(nodes[candidate] - position_m))
    if abs(nodes[index] - position_m) > 1.0e-12:
        raise ValueError(f"freeze reference lacks official position {position_m} m")
    return index


def _time_index(times: Iterable[float], time_s: int) -> int:
    times = tuple(times)
    index = min(range(len(times)), key=lambda candidate: abs(times[candidate] - time_s))
    if abs(times[index] - time_s) > 1.0e-9:
        raise ValueError(f"freeze reference lacks official time {time_s} s")
    return index


def _raw_value(payload: Dict[str, Any], field: str, time_s: int, position_m: float) -> float:
    ti = _time_index(payload["times_s"], time_s)
    pi = _position_index(payload["grid"]["nodes_m"], position_m)
    if field == "temperature_c":
        return float(payload["temperatures_k"][ti][pi]) - 273.15
    return float(payload["moistures_kg_kg"][ti][pi])


def _copy_and_fill(payload: Dict[str, Any], reuse_existing: bool = False) -> None:
    assert_not_official_output(CANDIDATE)
    CANDIDATE.parent.mkdir(parents=True, exist_ok=True)
    if CANDIDATE.exists():
        if reuse_existing:
            return
        raise FileExistsError(f"refusing to overwrite existing candidate: {CANDIDATE}")
    shutil.copy2(OFFICIAL_TEMPLATE, CANDIDATE)
    workbook = openpyxl.load_workbook(CANDIDATE)
    try:
        for sheet_name, field in EXPECTED_SHEETS:
            sheet = workbook[sheet_name]
            sheet.cell(1, 1).value = sheet.cell(1, 1).value
            for column, position_m in enumerate(OUTPUT_POSITIONS_M, start=2):
                sheet.cell(1, column).value = Decimal(str(position_m * 100.0)).quantize(Decimal("0.0"))
            for row, time_s in enumerate(range(1, 1801), start=2):
                sheet.cell(row, 1).value = time_s
                for column, position_m in enumerate(OUTPUT_POSITIONS_M, start=2):
                    sheet.cell(row, column).value = _rounded(_raw_value(payload, field, time_s, position_m))
        workbook.save(CANDIDATE)
    finally:
        workbook.close()


def _load_candidate_values() -> Dict[str, Tuple[Tuple[Any, ...], ...]]:
    workbook = openpyxl.load_workbook(CANDIDATE, read_only=True, data_only=False)
    try:
        return {
            sheet_name: tuple(tuple(row) for row in workbook[sheet_name].iter_rows(values_only=True))
            for sheet_name, _ in EXPECTED_SHEETS
        }
    finally:
        workbook.close()


def _trace(payload: Dict[str, Any], freeze_metrics: Dict[str, Any]) -> Dict[str, Any]:
    rng = random.Random(20260911)
    candidate_values = _load_candidate_values()

    def candidate_value(sheet_name: str, row: int, column: int) -> Any:
        return candidate_values[sheet_name][row - 1][column - 1]

    cells = []
    for sample_index in rng.sample(range(1800 * 21), 20):
        time_offset, position_index = divmod(sample_index, 21)
        time_s = time_offset + 1
        position_m = OUTPUT_POSITIONS_M[position_index]
        sheet_name, field = EXPECTED_SHEETS[sample_index % len(EXPECTED_SHEETS)]
        row = time_offset + 2
        column = position_index + 2
        raw_value = _raw_value(payload, field, time_s, position_m)
        candidate_cell = candidate_value(sheet_name, row, column)
        expected_value = _rounded(raw_value)
        cells.append({
            "sheet": sheet_name,
            "field": field,
            "row": row,
            "column": column,
            "time_s": time_s,
            "distance_cm": position_m * 100.0,
            "freeze_raw_value": raw_value,
            "candidate_value": candidate_cell,
            "expected_round_half_up": expected_value,
            "pass": candidate_cell == expected_value,
        })

    paper = {}
    for sheet_name, field in EXPECTED_SHEETS:
        entries = []
        for time_s in PAPER_TIMES_S:
            for position_m in PAPER_POSITIONS_M:
                position_index = round(position_m / 0.001)
                row = time_s + 1
                column = position_index + 2
                raw_value = _raw_value(payload, field, time_s, position_m)
                candidate_cell = candidate_value(sheet_name, row, column)
                expected_value = _rounded(raw_value)
                entries.append({
                    "time_s": time_s,
                    "distance_cm": position_m * 100.0,
                    "freeze_raw_value": raw_value,
                    "candidate_value": candidate_cell,
                    "expected_round_half_up": expected_value,
                    "pass": candidate_cell == expected_value,
                })
        paper[field] = {
            "count": len(entries),
            "pass_count": sum(item["pass"] for item in entries),
            "trace": entries,
        }
    return {
        "status": "PASS" if all(item["pass"] for item in cells) and all(item["pass_count"] == 35 for item in paper.values()) else "FAIL",
        "source": {
            "freeze_experiment": "experiments/Q1_FREEZE_RUN",
            "reference_path": str(SOURCE_REFERENCE.relative_to(ROOT)),
            "reference_sha256": sha256(SOURCE_REFERENCE),
            "official_template_path": str(OFFICIAL_TEMPLATE.relative_to(ROOT)),
            "official_template_sha256": sha256(OFFICIAL_TEMPLATE),
        },
        "random_seed": 20260911,
        "random_cell_trace_count": len(cells),
        "random_cell_trace": cells,
        "paper_tables": paper,
        "candidate_sha256": sha256(CANDIDATE),
        "freeze_metrics_reference": {
            "production_config_status": freeze_metrics.get("production_config_status"),
            "code_commit": freeze_metrics.get("code_commit"),
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reuse-existing", action="store_true", help="audit an already-generated candidate without overwriting it")
    args = parser.parse_args()
    if not SOURCE_REFERENCE.is_file():
        raise SystemExit(f"missing freeze source: {SOURCE_REFERENCE}")
    freeze_metrics_path = FREEZE_DIR / "metrics.json"
    freeze_metrics = json.loads(freeze_metrics_path.read_text(encoding="utf-8"))
    if freeze_metrics.get("status") != "COMPLETED" or freeze_metrics.get("production_config_status") != "Q1_PRODUCTION_CONFIG_FROZEN":
        raise SystemExit("Q1_FREEZE_RUN is not complete; candidate generation is blocked")
    if freeze_metrics.get("determinism", {}).get("status") != "PASS":
        raise SystemExit("Q1_FREEZE_RUN determinism check did not pass")

    spatial_metrics = json.loads((ROOT / "experiments" / "EXP-Q1-FULL-SPATIAL" / "metrics.json").read_text(encoding="utf-8"))
    temporal_metrics = json.loads((ROOT / "experiments" / "EXP-Q1-FULL-TEMPORAL" / "metrics.json").read_text(encoding="utf-8"))
    for label, metrics in (("spatial", spatial_metrics), ("temporal", temporal_metrics)):
        if metrics.get("criterion_reference", {}).get("pass") is not True:
            raise SystemExit(f"{label} full-horizon accuracy criterion did not pass")
        if metrics.get("selected_lowest_cost_passing_candidate") is None:
            raise SystemExit(f"{label} full-horizon production candidate is absent")

    payload = _load_payload(SOURCE_REFERENCE)
    _copy_and_fill(payload, reuse_existing=args.reuse_existing)
    trace = _trace(payload, freeze_metrics)
    trace_path = FREEZE_DIR / "candidate_trace.json"
    trace_path.write_text(json.dumps(trace, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    validator_errors = validate_candidate(ROOT, CANDIDATE)
    validation = {
        "status": "PASS" if trace["status"] == "PASS" and not validator_errors else "FAIL",
        "candidate_validator_errors": validator_errors,
        "trace_status": trace["status"],
        "random_cell_trace_count": trace["random_cell_trace_count"],
        "paper_table_trace": {
            field: {"count": values["count"], "pass_count": values["pass_count"]}
            for field, values in trace["paper_tables"].items()
        },
    }
    (FREEZE_DIR / "candidate_validation.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "status": validation["status"],
        "candidate": str(CANDIDATE.relative_to(ROOT)),
        "candidate_sha256": trace["candidate_sha256"],
        "random_cell_trace_count": trace["random_cell_trace_count"],
        "paper_table_trace": validation["paper_table_trace"],
        "validator_errors": validator_errors,
    }, ensure_ascii=False, indent=2))
    return 0 if validation["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
