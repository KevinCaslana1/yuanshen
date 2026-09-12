"""Fallback streaming author for the very large Q2 V3 workbook.

The artifact-tool object model was attempted first but cannot hold this
10-million-cell workbook within the available V8 heap.  This path uses
openpyxl's write_only workbook, preserves the official template's sheet names
and A1 labels, and streams the canonical Run 1 CSV exactly once.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import sys
import tempfile
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from openpyxl.cell import WriteOnlyCell

ROOT = Path(__file__).resolve().parents[1]
FREEZE = ROOT / "experiments" / "Q2_FREEZE_RUN_V3"
TEMPLATE = ROOT / "A题" / "附件" / "附件3" / "result2.xlsx"
CANDIDATE = ROOT / "deliverables" / "candidate" / "result2.xlsx"
SOURCE = FREEZE / "run_1" / "official_samples.csv"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def q4(raw: str) -> float:
    return float(Decimal(raw).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP))


def styled_cell(ws, value, number_format: str | None = None):
    cell = WriteOnlyCell(ws, value=value)
    if number_format is not None:
        cell.number_format = number_format
    return cell


def main() -> int:
    config = json.loads((FREEZE / "config.json").read_text(encoding="utf-8"))
    validation = json.loads((FREEZE / "validation.json").read_text(encoding="utf-8"))
    hashes = json.loads((FREEZE / "output_hashes.json").read_text(encoding="utf-8"))
    if validation.get("status") != "PRE_CANDIDATE_GATES_PASS" or not all(validation.get("gates", {}).values()):
        raise RuntimeError("pre-candidate gates are not all PASS")
    if sha256(SOURCE) != hashes["run_1"]["sampled_output"]:
        raise RuntimeError("V3 Run 1 sampled source hash mismatch")
    template_hash = sha256(TEMPLATE)
    if CANDIDATE.exists() and sha256(CANDIDATE) != template_hash:
        raise RuntimeError("refusing to overwrite a non-template candidate")
    template = openpyxl.load_workbook(TEMPLATE, read_only=True, data_only=False)
    try:
        sheet_names = tuple(template.sheetnames)
        a1_labels = {name: template[name]["A1"].value for name in sheet_names}
    finally:
        template.close()
    if len(sheet_names) != 2:
        raise RuntimeError(f"official template has unexpected sheets: {sheet_names!r}")

    wb = openpyxl.Workbook(write_only=True)
    ws_t = wb.create_sheet(sheet_names[0])
    ws_c = wb.create_sheet(sheet_names[1])
    for ws in (ws_t, ws_c):
        ws.append([styled_cell(ws, a1_labels[ws.title]), *[styled_cell(ws, index / 10, "0.0") for index in range(21)]])

    current_time: int | None = None
    temperature: list[float | None] = []
    moisture: list[float | None] = []
    rows_written = 0
    with SOURCE.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        expected_header = ["time_s", "radius_cm", "temperature_K", "temperature_C", "moisture_kg_kg"]
        if list(reader.fieldnames or []) != expected_header:
            raise RuntimeError("canonical source header mismatch")
        for row in reader:
            time_s = int(float(row["time_s"]))
            radius_index = int(round(float(row["radius_cm"]) * 10))
            if current_time is None:
                current_time = time_s
                temperature = [None] * 21
                moisture = [None] * 21
            if time_s != current_time:
                if len(temperature) != 21 or any(value is None for value in temperature + moisture):
                    raise RuntimeError(f"incomplete source layer at t={current_time}")
                ws_t.append([styled_cell(ws_t, current_time, "0"), *[styled_cell(ws_t, value, "0.0000") for value in temperature]])
                ws_c.append([styled_cell(ws_c, current_time, "0"), *[styled_cell(ws_c, value, "0.0000") for value in moisture]])
                rows_written += 1
                if time_s != current_time + 1:
                    raise RuntimeError(f"nonconsecutive source time {current_time}->{time_s}")
                current_time = time_s
                temperature = [None] * 21
                moisture = [None] * 21
            if not (0 <= radius_index <= 20) or temperature[radius_index] is not None:
                raise RuntimeError(f"duplicate/out-of-range source key t={time_s}, r={radius_index}")
            temperature[radius_index] = q4(row["temperature_C"])
            moisture[radius_index] = q4(row["moisture_kg_kg"])
    if current_time is not None:
        if any(value is None for value in temperature + moisture):
            raise RuntimeError(f"incomplete source layer at t={current_time}")
        ws_t.append([styled_cell(ws_t, current_time, "0"), *[styled_cell(ws_t, value, "0.0000") for value in temperature]])
        ws_c.append([styled_cell(ws_c, current_time, "0"), *[styled_cell(ws_c, value, "0.0000") for value in moisture]])
        rows_written += 1
    expected_horizon = int(config["final_horizon_s"])
    if rows_written != expected_horizon:
        raise RuntimeError(f"wrote {rows_written} rows, expected {expected_horizon}")
    for ws in (ws_t, ws_c):
        ws.freeze_panes = "B2"
        ws.sheet_view.showGridLines = True
    CANDIDATE.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix="result2_v3_", suffix=".xlsx", dir=CANDIDATE.parent)
    os.close(fd)
    temporary_path = Path(temporary)
    try:
        wb.save(temporary_path)
        os.replace(temporary_path, CANDIDATE)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
    print(json.dumps({"status": "PASS", "authoring_mode": "openpyxl_write_only_streaming_fallback", "candidate": str(CANDIDATE), "source_sha256": sha256(SOURCE), "template_sha256": template_hash, "rows_written": rows_written, "candidate_sha256": sha256(CANDIDATE)}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"{type(exc).__name__}: {exc}", file=sys.stderr)
        raise
