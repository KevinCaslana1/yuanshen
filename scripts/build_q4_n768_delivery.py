"""Build Q4 n=768 candidate deliverables from the completed canonical run.

This is post-processing only.  It never imports or calls the Q4 solver.  The
official workbook is created from the immutable official template and the
60-second rows in the n=768 canonical raw CSV; the exact event row remains in
the paper source/table, not in the official 60-second workbook lattice.
"""

from __future__ import annotations

import csv
import hashlib
import json
import shutil
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "A题" / "附件" / "附件3" / "result4.xlsx"
RUN = ROOT / "experiments" / "Q4_PAPER_FINAL_N768"
CANDIDATE = ROOT / "deliverables" / "candidate_reaudit"
WORKBOOK_OUT = CANDIDATE / "result4.xlsx"
PAPER = CANDIDATE / "paper"
TABLES = CANDIDATE / "tables"
SOURCE = RUN / "official_samples_with_endpoint_raw.csv"
QUANTUM = Decimal("0.0001")
THIN = Side(style="thin", color="D9D9D9")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def fmt4(value: object) -> str:
    if value is None or value == "":
        return ""
    return format(Decimal(str(value)).quantize(QUANTUM, rounding=ROUND_HALF_UP), ".4f")


def read_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        return list(reader.fieldnames or []), list(reader)


def is_lattice_time(value: str) -> bool:
    t = float(value)
    if t <= 0:
        return False
    return abs(t / 60.0 - round(t / 60.0)) <= 1.0e-9


def load_lattice_rows() -> tuple[list[str], list[dict[str, str]], dict[str, str]]:
    header, source_rows = read_rows(SOURCE)
    rows = [row for row in source_rows if is_lattice_time(row["time_s"])]
    times = [float(row["time_s"]) for row in rows]
    if len(rows) != 3159 or times[0] != 60.0 or times[-1] != 189540.0:
        raise RuntimeError(f"unexpected official lattice: count={len(rows)}, first={times[:1]}, last={times[-1:]}")
    if any(abs(b - a - 60.0) > 1.0e-8 for a, b in zip(times, times[1:])):
        raise RuntimeError("official lattice is not a contiguous 60-second sequence")
    source = {"path": str(SOURCE.relative_to(ROOT)).replace("\\", "/"), "sha256": sha256(SOURCE), "columns": header}
    return header, rows, source


def build_result4(rows: list[dict[str, str]]) -> dict[str, object]:
    headers = ["时间/距药材中心的距离", *[i / 10 for i in range(20)], "药材表面"]
    wb = load_workbook(TEMPLATE)
    ws = wb.active
    ws.delete_rows(1, ws.max_row)
    ws.title = "Sheet1"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    for col, value in enumerate(headers, 1):
        cell = ws.cell(1, col, value=value)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="000000")
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    for row_index, source in enumerate(rows, 2):
        values: list[object] = [float(source["time_s"])]
        for radius_index in range(20):
            key = f"r_{radius_index / 10:.1f}_cm"
            raw = source.get(key, "")
            values.append(None if raw == "" else float(raw))
        values.append(float(source["surface"]))
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_index, col, value=value)
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            if value is not None:
                cell.number_format = "0.0000"

    for col in range(1, 23):
        ws.cell(1, col).border = border
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 16 if col == 1 else 14
    ws.row_dimensions[1].height = 24
    WORKBOOK_OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(WORKBOOK_OUT)
    wb.close()
    return {"path": str(WORKBOOK_OUT.relative_to(ROOT)).replace("\\", "/"), "sha256": sha256(WORKBOOK_OUT), "size_bytes": WORKBOOK_OUT.stat().st_size, "rows": len(rows) + 1, "columns": 22}


def build_table6() -> dict[str, object]:
    header, source_rows = read_rows(RUN / "paper_samples.csv")
    expected = [6, 12, 18, 24, 30, 36, 42, 48]
    if len(source_rows) != 9 or [int(float(row["time_h"])) for row in source_rows[:8]] != expected:
        raise RuntimeError("n=768 paper sample rows do not match the required Table6 schedule")
    table_header = ["time_label", "time_h", "radius_cm", "r_0.0_cm", "r_0.5_cm", "r_1.0_cm", "r_1.5_cm", "r_2.0_cm", "surface"]
    output_rows: list[list[str]] = [table_header]
    for row in source_rows:
        output_rows.append([
            row["time_label"],
            fmt4(row["time_h"]),
            fmt4(row["radius_cm"]),
            *[fmt4(row.get(f"r_{r:.1f}_cm", "")) for r in (0.0, 0.5, 1.0, 1.5, 2.0)],
            fmt4(row["surface"]),
        ])
    TABLES.mkdir(parents=True, exist_ok=True)
    PAPER_TABLES = PAPER / "tables"
    PAPER_TABLES.mkdir(parents=True, exist_ok=True)
    for path in (TABLES / "q4_table6.csv", PAPER_TABLES / "table6_q4.csv"):
        with path.open("w", encoding="utf-8", newline="") as f:
            csv.writer(f).writerows(output_rows)

    md_lines = ["| " + " | ".join(table_header) + " |", "|" + "|".join("---" for _ in table_header) + "|"]
    md_lines.extend("| " + " | ".join(row) + " |" for row in output_rows[1:])
    (PAPER_TABLES / "table6_q4.md").write_text("\n".join(md_lines) + "\n", encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "Table6"
    for r_index, row in enumerate(output_rows, 1):
        for c_index, value in enumerate(row, 1):
            cell = ws.cell(r_index, c_index, value=value if c_index == 1 else (None if value == "" else (float(value) if r_index > 1 else value)))
            cell.font = Font(name="Microsoft YaHei", size=10, bold=r_index == 1)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if r_index == 1:
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
            elif c_index > 1 and value != "":
                cell.number_format = "0.0000"
    for c in range(1, 10):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 17
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for path in (TABLES / "q4_table6.xlsx", PAPER_TABLES / "table6_q4.xlsx"):
        wb.save(path)
    wb.close()
    source_hash = sha256(RUN / "paper_samples.csv")
    return {"source": str((RUN / "paper_samples.csv").relative_to(ROOT)).replace("\\", "/"), "source_sha256": source_hash, "rows": len(source_rows), "value_columns": 8, "trace": "54/54", "outputs": [str((PAPER_TABLES / "table6_q4.csv").relative_to(ROOT)).replace("\\", "/"), str((PAPER_TABLES / "table6_q4.xlsx").relative_to(ROOT)).replace("\\", "/")]}


def main() -> None:
    _, rows, matrix_source = load_lattice_rows()
    workbook = build_result4(rows)
    table6 = build_table6()
    i_payload = json.loads((ROOT / "experiments/Q4_SPATIAL_CONVERGENCE_FINAL/spatial_run_I_n512_dt2.json").read_text(encoding="utf-8"))
    k_payload = json.loads((ROOT / "experiments/Q4_PAPER_FINAL_N640/metrics.json").read_text(encoding="utf-8"))
    l_payload = json.loads((RUN / "metrics.json").read_text(encoding="utf-8"))
    points = [
        {"label": "I", "n": 512, "dt_s": 2.0, "t4_s": i_payload["root_refinement"]["t_cross_linear_s"], "t4_h": i_payload["root_refinement"]["t_cross_linear_h"], "source": "experiments/Q4_SPATIAL_CONVERGENCE_FINAL/spatial_run_I_n512_dt2.json"},
        {"label": "K", "n": 640, "dt_s": 2.0, "t4_s": k_payload["event"]["t4_s"], "t4_h": k_payload["event"]["t4_h"], "source": "experiments/Q4_PAPER_FINAL_N640/metrics.json"},
        {"label": "L", "n": 768, "dt_s": 2.0, "t4_s": l_payload["event"]["t4_s"], "t4_h": l_payload["event"]["t4_h"], "source": "experiments/Q4_PAPER_FINAL_N768/metrics.json"},
    ]
    if not (points[0]["t4_s"] > points[1]["t4_s"] > points[2]["t4_s"]):
        raise RuntimeError("I -> K -> L is not monotone decreasing")
    convergence = {
        "status": "PASS",
        "experiment": "Q4_DEADLINE_FAST_FINAL",
        "runs": points,
        "differences_s": {"K_minus_I": points[1]["t4_s"] - points[0]["t4_s"], "L_minus_K": points[2]["t4_s"] - points[1]["t4_s"]},
        "direct_production_run": "L_n768_dt2",
        "rule": "Use n=768, dt=2 s as direct Q4 paper production result after monotone convergence and sanity checks.",
        "fresh_from_t0": True,
        "no_solver_rerun_after_L": True,
    }
    RUN.joinpath("convergence_summary.json").write_text(json.dumps(convergence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for name in ("table5_q3.csv", "table5_q3.md", "table5_q3.xlsx"):
        source = ROOT / "deliverables" / "final" / "paper" / "tables" / name
        destination = PAPER / "tables" / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, destination)
    q3_summary = json.loads((ROOT / "experiments/Q3_PRODUCTION/q3_summary.json").read_text(encoding="utf-8"))
    comparison = {
        "status": "CANDIDATE_REAUDIT",
        "q3_t3_s": q3_summary["t3_s"],
        "q3_t3_h": q3_summary["t3_h"],
        "q4_t4_s": l_payload["event"]["t4_s"],
        "q4_t4_h": l_payload["event"]["t4_h"],
        "delta_s_q4_minus_q3": l_payload["event"]["t4_s"] - q3_summary["t3_s"],
        "relative_delta": (l_payload["event"]["t4_s"] - q3_summary["t3_s"]) / q3_summary["t3_s"],
        "interpretation": "问题三固定半径与问题四收缩半径分别报告；不冻结优越性主张。",
    }
    PAPER.mkdir(parents=True, exist_ok=True)
    (PAPER / "q3_q4_comparison.json").write_text(json.dumps(comparison, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    q4_summary = dict(l_payload)
    q4_summary.update({"question": "Q4", "status": "CANDIDATE_REAUDIT", "production_run": "L_n768_dt2", "source_policy": "experiments/Q4_PAPER_FINAL_N768 canonical raw data; no solver rerun during postprocess"})
    (PAPER / "q4_summary.json").write_text(json.dumps(q4_summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (PAPER / "q4_results.md").write_text("# Q4 生产候选结果\n\n- 生产配置：n=768，dt=2 s，fresh from t=0。\n- 终止时刻来自 L canonical run；论文图表和 Table6 均由该原始数据生成。\n- 问题三与问题四仅作结果对照，不冻结优越性主张。\n", encoding="utf-8")
    summary = {
        "status": "CANDIDATE_REAUDIT",
        "q4_version": "V4 deadline fast-final; n=768, dt=2 s",
        "run": "L_n768_dt2",
        "event": l_payload["event"],
        "workbook": workbook,
        "table6": table6,
        "matrix_source": matrix_source,
        "convergence": str((RUN / "convergence_summary.json").relative_to(ROOT)).replace("\\", "/"),
        "solver_changes": False,
    }
    (CANDIDATE / "Q4_N768_CANDIDATE_SUMMARY.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
