"""Validate the approved Q1 final workbook and its candidate byte identity."""

from __future__ import annotations

import argparse
import json
import math
import sys
from decimal import Decimal
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - environment failure
    raise SystemExit("openpyxl is required for Q1 final validation") from exc

from scripts.validate_inputs import EXPECTED, sha256
from scripts.validate_q1_candidate import _decimal, _reported_value_ok, _same_decimal


EXPECTED_SHEETS = ("温度", "水分浓度")
EXPECTED_ROWS = 1801
EXPECTED_COLUMNS = 22
EXPECTED_TIMES = tuple(range(1, 1801))
EXPECTED_DISTANCES = tuple(Decimal(index) / Decimal(10) for index in range(21))


def _within(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False


def _formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _validate_sheet(
    workbook: Any,
    template: Any,
    sheet_name: str,
    errors: list[str],
) -> tuple[tuple[Decimal, ...], tuple[Decimal, ...]]:
    sheet = workbook[sheet_name]
    if sheet.max_row != EXPECTED_ROWS or sheet.max_column != EXPECTED_COLUMNS:
        errors.append(
            f"{sheet_name} shape mismatch: expected={EXPECTED_ROWS}x{EXPECTED_COLUMNS} "
            f"actual={sheet.max_row}x{sheet.max_column}"
        )

    if sheet["A1"].value != template[sheet_name]["A1"].value:
        errors.append(f"{sheet_name}!A1 does not match the official template header")

    rows = sheet.iter_rows(
        min_row=1,
        max_row=EXPECTED_ROWS,
        min_col=1,
        max_col=EXPECTED_COLUMNS,
        values_only=True,
    )
    header = tuple(next(rows, ()))
    distances: list[Decimal] = []
    for column, expected_distance in enumerate(EXPECTED_DISTANCES, start=2):
        value = header[column - 1] if len(header) >= column else None
        if not _same_decimal(value, expected_distance):
            errors.append(
                f"{sheet_name}!{get_column_letter(column)}1 distance header is not {expected_distance} cm"
            )
        decimal_value = _decimal(value)
        if decimal_value is not None:
            distances.append(decimal_value)

    times: list[Decimal] = []
    for row_number, expected_time in enumerate(EXPECTED_TIMES, start=2):
        values = tuple(next(rows, ()))
        time_value = values[0] if values else None
        if not _same_decimal(time_value, Decimal(expected_time)):
            errors.append(f"{sheet_name}!A{row_number} is not the strict 1..1800 time sequence")
        decimal_time = _decimal(time_value)
        if decimal_time is not None:
            times.append(decimal_time)

        for column in range(1, EXPECTED_COLUMNS + 1):
            value = values[column - 1] if len(values) >= column else None
            if _formula(value):
                errors.append(f"{sheet_name}!{get_column_letter(column)}{row_number} contains a formula")
            if column >= 2 and not _reported_value_ok(value):
                errors.append(
                    f"{sheet_name}!{get_column_letter(column)}{row_number} is not a finite "
                    "numeric value already rounded to 4 decimals"
                )
                if len(errors) >= 100:
                    errors.append("error list truncated after 100 issues")
                    return tuple(times), tuple(distances)

    return tuple(times), tuple(distances)


def validate_final(root: Path, final_path: Path | None = None, candidate_path: Path | None = None) -> dict[str, Any]:
    """Return a machine-readable final-workbook validation report."""

    errors: list[str] = []
    final_root = (root / "deliverables" / "final").resolve()
    candidate_root = (root / "deliverables" / "candidate").resolve()
    official_root = (root / "A题").resolve()
    final = (final_path or final_root / "result1.xlsx").resolve()
    candidate = (candidate_path or candidate_root / "result1.xlsx").resolve()
    official_template = root / "A题" / "附件" / "附件3" / "result1.xlsx"

    if not _within(final, final_root):
        errors.append(f"final path escapes deliverables/final: {final}")
    if _within(final, official_root):
        errors.append("final path overlaps A题 official source")
    if not final.is_file():
        errors.append(f"missing Q1 final workbook: {final}")
    if not candidate.is_file():
        errors.append(f"missing Q1 candidate workbook: {candidate}")
    if not official_template.is_file():
        errors.append(f"missing official Q1 template: {official_template}")

    candidate_hash = sha256(candidate) if candidate.is_file() else None
    final_hash = sha256(final) if final.is_file() else None
    template_hash = sha256(official_template) if official_template.is_file() else None
    expected_template_hash = EXPECTED["A题/附件/附件3/result1.xlsx"]["sha256"]
    if template_hash != expected_template_hash:
        errors.append(f"official Q1 template hash changed: expected={expected_template_hash} actual={template_hash}")
    if candidate_hash and final_hash and candidate_hash != final_hash:
        errors.append(f"candidate/final SHA-256 mismatch: candidate={candidate_hash} final={final_hash}")

    sheet_reports: dict[str, Any] = {}
    if final.is_file() and official_template.is_file():
        try:
            workbook = openpyxl.load_workbook(final, read_only=True, data_only=False)
            template = openpyxl.load_workbook(official_template, read_only=True, data_only=False)
        except Exception as exc:
            errors.append(f"cannot open final or official workbook: {exc}")
        else:
            try:
                actual_sheets = tuple(workbook.sheetnames)
                if actual_sheets != EXPECTED_SHEETS:
                    errors.append(f"sheet mismatch: expected={EXPECTED_SHEETS} actual={actual_sheets}")
                axes: dict[str, tuple[tuple[Decimal, ...], tuple[Decimal, ...]]] = {}
                for sheet_name in EXPECTED_SHEETS:
                    if sheet_name not in actual_sheets:
                        errors.append(f"missing required sheet: {sheet_name}")
                        continue
                    times, distances = _validate_sheet(workbook, template, sheet_name, errors)
                    axes[sheet_name] = (times, distances)
                    sheet_reports[sheet_name] = {
                        "rows": len(times),
                        "columns": len(distances) + 1,
                        "finite_values": True,
                        "four_decimal_values": True,
                    }
                if len(axes) == 2:
                    if axes["温度"][0] != axes["水分浓度"][0]:
                        errors.append("temperature/moisture time axes differ")
                    if axes["温度"][1] != axes["水分浓度"][1]:
                        errors.append("temperature/moisture spatial axes differ")
            finally:
                template.close()
                workbook.close()

    checks = {
        "final_exists": final.is_file(),
        "candidate_exists": candidate.is_file(),
        "official_template_hash_unchanged": template_hash == expected_template_hash,
        "candidate_final_sha256_equal": bool(candidate_hash and final_hash and candidate_hash == final_hash),
        "required_sheets_and_shape": not any("sheet mismatch" in error or "shape mismatch" in error for error in errors),
        "strict_time_and_space_axes": not any("sequence" in error or "distance header" in error or "axes differ" in error for error in errors),
        "finite_four_decimal_values": not any("finite" in error or "formula" in error for error in errors),
        "no_extra_sheets": not any("sheet mismatch" in error for error in errors),
    }
    status = "PASS" if not errors and all(checks.values()) else "FAIL"
    return {
        "status": status,
        "validation_id": "FINAL_RESULT1_VALIDATION",
        "final_path": str(final.relative_to(root)) if _within(final, root) else str(final),
        "candidate_path": str(candidate.relative_to(root)) if _within(candidate, root) else str(candidate),
        "candidate_sha256": candidate_hash,
        "final_sha256": final_hash,
        "official_template_sha256": template_hash,
        "expected_template_sha256": expected_template_hash,
        "expected_shape": {"rows": EXPECTED_ROWS, "columns": EXPECTED_COLUMNS},
        "sheet_reports": sheet_reports,
        "checks": checks,
        "errors": errors,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--output", type=Path, default=None)
    args = parser.parse_args()
    root = args.root.resolve()
    report = validate_final(root)
    if args.output:
        output = args.output if args.output.is_absolute() else root / args.output
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if report["status"] == "PASS":
        print("FINAL_RESULT1_VALIDATION = PASS")
        print(f"candidate/final SHA-256 = {report['candidate_sha256']}")
        print("sheets, axes, finite values, formulas, and 4-decimal output validated")
        return 0
    for error in report["errors"]:
        print(f"[FAIL] {error}")
    print(f"FINAL_RESULT1_VALIDATION = FAIL ({len(report['errors'])} issue(s))")
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
