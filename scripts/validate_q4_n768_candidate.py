"""Independent audit of the Q4 n=768 candidate package.

The validator reads only completed outputs and frozen sources.  It does not
invoke either Q3 or Q4 numerical code.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
CAND = ROOT / "deliverables" / "candidate_reaudit"
RUN = ROOT / "experiments" / "Q4_PAPER_FINAL_N768"
Q4_MATRIX = RUN / "official_samples_with_endpoint_raw.csv"
PAPER_SAMPLES = RUN / "paper_samples.csv"
FIG_TRACE = CAND / "paper" / "Q4_N768_FIGURE_TRACE.json"
TABLE_TRACE = CAND / "paper" / "TABLE_1_6_TRACE.json"
OLD_TABLE_TRACE = ROOT / "deliverables" / "final" / "paper" / "TABLE_1_6_TRACE.json"
QUANTUM = Decimal("0.0001")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def fmt4(value: object) -> str:
    if value is None or value == "":
        return ""
    return format(Decimal(str(value)).quantize(QUANTUM, rounding=ROUND_HALF_UP), ".4f")


def csv_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        return list(reader.fieldnames or []), list(reader)


def array_hash(*arrays: list[float]) -> str:
    h = hashlib.sha256()
    for values in arrays:
        h.update(__import__("numpy").asarray(values, dtype="<f8").tobytes(order="C"))
    return h.hexdigest()


def validate_workbook() -> dict[str, object]:
    workbook_path = CAND / "result4.xlsx"
    source_header, source_rows = csv_rows(Q4_MATRIX)
    lattice = [row for row in source_rows if float(row["time_s"]) > 0 and abs(float(row["time_s"]) / 60 - round(float(row["time_s"]) / 60)) <= 1e-9]
    wb = load_workbook(workbook_path, read_only=False, data_only=False)
    ws = wb.active
    expected_header = ["时间/距药材中心的距离", *[i / 10 for i in range(20)], "药材表面"]
    header_ok = [ws.cell(1, c).value for c in range(1, 23)] == expected_header
    formulas = []
    nonfinite = []
    format_bad = []
    trace_total = 0
    trace_pass = 0
    for r_index, source in enumerate(lattice, 2):
        expected = [float(source["time_s"])]
        expected.extend(None if source.get(f"r_{i / 10:.1f}_cm", "") == "" else float(source[f"r_{i / 10:.1f}_cm"]) for i in range(20))
        expected.append(float(source["surface"]))
        for c_index, expected_value in enumerate(expected, 1):
            cell = ws.cell(r_index, c_index)
            value = cell.value
            if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                formulas.append(cell.coordinate)
            if value is not None and isinstance(value, (int, float)) and not math.isfinite(float(value)):
                nonfinite.append(cell.coordinate)
            if expected_value is None:
                if value is not None:
                    trace_total += 1
                continue
            trace_total += 1
            if value == expected_value or (isinstance(value, (int, float)) and math.isclose(float(value), expected_value, rel_tol=0.0, abs_tol=5.0e-15)):
                trace_pass += 1
            else:
                raise AssertionError(f"workbook trace mismatch {cell.coordinate}: {value!r} != {expected_value!r}")
            if cell.number_format != "0.0000":
                format_bad.append(cell.coordinate)
    row_times = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    time_ok = len(row_times) == len(lattice) and row_times[0] == 60.0 and row_times[-1] == 189540.0 and all(abs(float(b) - float(a) - 60.0) < 1e-9 for a, b in zip(row_times, row_times[1:]))
    shape_ok = ws.max_row == 3160 and ws.max_column == 22
    result = {"path": str(workbook_path.relative_to(ROOT)).replace("\\", "/"), "sha256": sha256(workbook_path), "size_bytes": workbook_path.stat().st_size, "shape": f"{ws.max_row}x{ws.max_column}", "shape_pass": shape_ok, "header_pass": header_ok, "time_lattice_pass": time_ok, "formula_count": len(formulas), "nonfinite_count": len(nonfinite), "number_format_bad_count": len(format_bad), "trace_all": f"{trace_pass}/{trace_total}", "trace_100": "100/100" if trace_pass >= 100 else f"{trace_pass}/100", "status": "PASS" if shape_ok and header_ok and time_ok and not formulas and not nonfinite and not format_bad and trace_pass == trace_total else "FAIL"}
    wb.close()
    if result["status"] != "PASS":
        raise AssertionError(result)
    return result


def validate_table6() -> dict[str, object]:
    _, source = csv_rows(PAPER_SAMPLES)
    _, target = csv_rows(CAND / "paper" / "tables" / "table6_q4.csv")
    if len(source) != len(target):
        raise AssertionError(f"Table6 row count mismatch {len(target)} != {len(source)}")
    numeric_keys = ["time_h", "radius_cm", "r_0.0_cm", "r_0.5_cm", "r_1.0_cm", "r_1.5_cm", "r_2.0_cm", "surface"]
    checked = 0
    for s, t in zip(source, target):
        if s["time_label"] != t["time_label"]:
            raise AssertionError("Table6 label mismatch")
        for key in numeric_keys:
            expected = fmt4(s.get(key, ""))
            if t[key] != expected:
                raise AssertionError(f"Table6 mismatch {key}: {t[key]!r} != {expected!r}")
            if expected != "":
                checked += 1
    wb = load_workbook(CAND / "paper" / "tables" / "table6_q4.xlsx", read_only=False, data_only=False)
    ws = wb.active
    xlsx_shape = (ws.max_row, ws.max_column)
    formula_count = sum(1 for row in ws.iter_rows() for cell in row if cell.data_type == "f")
    wb.close()
    return {"csv_rows": len(target), "valid_cells": f"{checked}/{checked}", "xlsx_shape": f"{xlsx_shape[0]}x{xlsx_shape[1]}", "formula_count": formula_count, "status": "PASS" if checked == 54 and xlsx_shape == (10, 9) and formula_count == 0 else "FAIL"}


def validate_figures() -> dict[str, object]:
    manifest = json.loads(FIG_TRACE.read_text(encoding="utf-8"))
    expected_cn = {
        "fig_5_15_q4_radius_pchip": ["问题四药材半径演化", "时间 t / h", "半径 R(t) / cm"],
        "fig_5_16_q4_dynamic_radial_moisture": ["问题四不同时间的径向水分浓度分布", "水分浓度 C / (kg/kg)"],
        "fig_5_17_q4_radius_and_mean": ["问题四半径与体积加权平均水分浓度", "体积加权平均值 C / (kg/kg)"],
        "fig_5_18_q3_q4_drying_time_comparison": ["问题三与问题四阈值时间对比", "问题三（固定半径）", "问题四（收缩半径）"],
    }
    actual = []
    for item in manifest["figures"]:
        stem = item["figure"]
        png = ROOT / item["output"]["png"]
        svg = ROOT / item["output"]["svg"]
        with Image.open(png) as image:
            dpi = image.info.get("dpi", (0.0, 0.0))
            dpi_pass = float(dpi[0]) >= 300.0 and float(dpi[1]) >= 300.0
            size_pass = image.size == (2160, 1380)
        text = svg.read_text(encoding="utf-8")
        text_pass = all(term in text for term in expected_cn[stem]) and "\ufffd" not in text and "□" not in text
        hash_pass = sha256(png) == item["output"]["png_sha256"] and sha256(svg) == item["output"]["svg_sha256"]
        actual.append({"figure": stem, "dpi": dpi, "dpi_pass": dpi_pass, "size_pass": size_pass, "chinese_text_pass": text_pass, "hash_pass": hash_pass, "trace": item["series"]})
    all_pass = len(actual) == 4 and all(row["dpi_pass"] and row["size_pass"] and row["chinese_text_pass"] and row["hash_pass"] for row in actual)
    return {"figure_groups": len(actual), "png": len(actual), "svg": len(actual), "dpi": f"{sum(row['dpi_pass'] for row in actual)}/{len(actual)}", "data_trace": manifest["counts"]["data_trace"], "chinese_text": f"{sum(row['chinese_text_pass'] for row in actual)}/{len(actual)}", "details": actual, "status": "PASS" if all_pass else "FAIL"}


def validate_frozen_tables() -> dict[str, object]:
    old = json.loads(OLD_TABLE_TRACE.read_text(encoding="utf-8"))
    new = json.loads(TABLE_TRACE.read_text(encoding="utf-8"))
    unchanged = old["tables"][:5] == new["tables"][:5]
    return {"table1_5_trace_unchanged": unchanged, "old_counts": {x["table"]: x["trace"] for x in old["tables"][:5]}, "new_counts": {x["table"]: x["trace"] for x in new["tables"][:5]}, "status": "PASS" if unchanged else "FAIL"}


def main() -> None:
    workbook = validate_workbook()
    table6 = validate_table6()
    figures = validate_figures()
    frozen_tables = validate_frozen_tables()
    docs = {"word": (CAND / "paper" / "论文表1-表6_可直接复制.docx").exists(), "pdf": (CAND / "paper" / "论文表1-表6_预览.pdf").exists()}
    report = {"status": "PASS" if all(x["status"] == "PASS" for x in (workbook, table6, figures, frozen_tables)) and all(docs.values()) else "FAIL", "workbook": workbook, "table6": table6, "figures": figures, "frozen_tables": frozen_tables, "merged_documents": docs}
    (CAND / "Q4_N768_VALIDATION.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))
    if report["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
