"""Read-only tail-window audit for official Attachment 1.

The output is an experiment record. The official XLSX is never modified.
"""

from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path
from statistics import mean, pstdev

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "A题" / "附件" / "附件1.xlsx"
OUT = ROOT / "experiments" / "EXP-Q2-ENV-TAIL" / "metrics.json"


def slope(xs: list[float], ys: list[float]) -> float:
    xbar = mean(xs)
    ybar = mean(ys)
    numerator = sum((x - xbar) * (y - ybar) for x, y in zip(xs, ys))
    denominator = sum((x - xbar) ** 2 for x in xs)
    return numerator / denominator


def summarize(rows: list[tuple[float, float, float]], n: int) -> dict[str, object]:
    tail = rows[-n:]
    times = [r[0] for r in tail]
    temperatures = [r[1] for r in tail]
    concentrations = [r[2] for r in tail]
    return {
        "n": n,
        "time_start_s": times[0],
        "time_end_s": times[-1],
        "temperature_C": {
            "mean": mean(temperatures),
            "population_std": pstdev(temperatures),
            "last": temperatures[-1],
            "min": min(temperatures),
            "max": max(temperatures),
            "linear_slope_per_s": slope(times, temperatures),
        },
        "moisture_kg_kg": {
            "mean": mean(concentrations),
            "population_std": pstdev(concentrations),
            "last": concentrations[-1],
            "min": min(concentrations),
            "max": max(concentrations),
            "linear_slope_per_s": slope(times, concentrations),
        },
    }


def main() -> None:
    digest = hashlib.sha256(INPUT.read_bytes()).hexdigest()
    workbook = load_workbook(INPUT, data_only=True, read_only=True)
    worksheet = workbook["Sheet1"]
    rows = [
        (float(row[0]), float(row[1]), float(row[2]))
        for row in worksheet.iter_rows(min_row=2, values_only=True)
    ]
    workbook.close()
    times = [r[0] for r in rows]
    payload = {
        "experiment": "EXP-Q2-ENV-TAIL",
        "status": "COMPLETED_DESIGN_STAGE_AUDIT",
        "input": "A题/附件/附件1.xlsx",
        "input_sha256": digest,
        "sheet": "Sheet1",
        "row_count": len(rows),
        "time_start_s": times[0],
        "time_end_s": times[-1],
        "constant_interval_s": sorted({b - a for a, b in zip(times, times[1:])}),
        "finite_values": all(math.isfinite(v) for row in rows for v in row),
        "tail_windows": [summarize(rows, n) for n in (10, 20, 40, 80)],
        "team_reference_candidate_after_14400s": {
            "temperature_C": 50.00,
            "moisture_kg_kg": 0.0500,
            "classification": "TEAM_REFERENCE_NUMERICAL_PROPOSAL_NOT_OFFICIAL",
        },
        "formal_Q2_solver_run": False,
        "workbook_written": False,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
