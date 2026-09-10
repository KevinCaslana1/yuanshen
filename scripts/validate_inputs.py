"""Validate official input assets without modifying or solving the problem.

This script checks existence, file size, SHA-256, workbook sheet names, and
basic row/column dimensions. It never writes to A题/ and never computes an
answer, calls a solver, or fills a result workbook.
"""

from __future__ import annotations

import argparse
import hashlib
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - environment failure
    raise SystemExit("openpyxl is required for read-only XLSX validation") from exc


EXPECTED = {
    "A题/A题.pdf": {
        "bytes": 553320,
        "sha256": "052d8014bff5727c019b72e44fdffaf5c145ce04050dd938baaf3527db331736",
    },
    "A题/附件/附件1.xlsx": {
        "bytes": 16586,
        "sha256": "7ef32870abeef420b89560b2530ff60dfe4255917805151d89988d0311af9dd7",
        "sheets": {"Sheet1": (242, 3)},
    },
    "A题/附件/附件2.xlsx": {
        "bytes": 11485,
        "sha256": "5563acbfa4b4afb10cc6c03e2207e5369bf39da27576672aff14cc5c32e704af",
        "sheets": {"Sheet1": (146, 2)},
    },
    "A题/附件/附件3/result1.xlsx": {
        "bytes": 10073,
        "sha256": "23b261b295c1b787d000eebbca6521c37075107b6fcf78724f8d395ce1798ff4",
        "sheets": {"温度": (5, 6), "水分浓度": (5, 6)},
    },
    "A题/附件/附件3/result2.xlsx": {
        "bytes": 10073,
        "sha256": "23b261b295c1b787d000eebbca6521c37075107b6fcf78724f8d395ce1798ff4",
        "sheets": {"温度": (5, 6), "水分浓度": (5, 6)},
    },
    "A题/附件/附件3/result3.xlsx": {
        "bytes": 9345,
        "sha256": "07e4793d620a7f899804c0298d49a16a197960440fd47f8bb780c57ec27e2859",
        "sheets": {"Sheet1": (5, 6)},
    },
    "A题/附件/附件3/result4.xlsx": {
        "bytes": 9351,
        "sha256": "86e9300ffa3d30c43de895ea6723da943e85b8740b137bcae5af7107f076eeac",
        "sheets": {"Sheet1": (5, 6)},
    },
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate(root: Path) -> int:
    failures = 0

    for relative, expected in EXPECTED.items():
        asset_failures = 0
        path = root / Path(relative)
        if not path.is_file():
            print(f"[FAIL] missing: {relative}")
            failures += 1
            continue

        actual_bytes = path.stat().st_size
        actual_hash = sha256(path)
        if actual_bytes != expected["bytes"]:
            print(f"[FAIL] size: {relative} expected={expected['bytes']} actual={actual_bytes}")
            failures += 1
            asset_failures += 1
        if actual_hash != expected["sha256"]:
            print(f"[FAIL] sha256: {relative} expected={expected['sha256']} actual={actual_hash}")
            failures += 1
            asset_failures += 1

        expected_sheets = expected.get("sheets")
        if expected_sheets:
            try:
                workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
                actual_sheets = {
                    worksheet.title: (worksheet.max_row, worksheet.max_column)
                    for worksheet in workbook.worksheets
                }
                workbook.close()
            except Exception as exc:
                print(f"[FAIL] workbook unreadable: {relative}: {exc}")
                failures += 1
                asset_failures += 1
            else:
                if actual_sheets != expected_sheets:
                    print(f"[FAIL] workbook structure: {relative} expected={expected_sheets} actual={actual_sheets}")
                    failures += 1
                    asset_failures += 1

        if asset_failures == 0:
            print(f"[PASS] {relative}")

    if failures:
        print(f"Input validation failed: {failures} issue(s)")
        return 1
    print(f"Input validation passed: {len(EXPECTED)} official asset(s)")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    return validate(args.root.resolve())


if __name__ == "__main__":
    sys.exit(main())
