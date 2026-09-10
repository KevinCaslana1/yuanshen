"""Validate official result-template structure without writing to the source files."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - environment failure
    raise SystemExit("openpyxl is required for read-only XLSX validation") from exc

try:
    from scripts.validate_inputs import EXPECTED, sha256
except ModuleNotFoundError:  # direct invocation: python scripts/validate_templates.py
    from validate_inputs import EXPECTED, sha256


TEMPLATES = {
    "A题/附件/附件3/result1.xlsx": {
        "sheets": ("温度", "水分浓度"),
        "time_examples": (1, 2, 3, "…"),
        "last_header": 2,
    },
    "A题/附件/附件3/result2.xlsx": {
        "sheets": ("温度", "水分浓度"),
        "time_examples": (1, 2, 3, "…"),
        "last_header": 2,
    },
    "A题/附件/附件3/result3.xlsx": {
        "sheets": ("Sheet1",),
        "time_examples": (60, 120, 180, "…"),
        "last_header": 2,
    },
    "A题/附件/附件3/result4.xlsx": {
        "sheets": ("Sheet1",),
        "time_examples": (60, 120, 180, "…"),
        "last_header": "药材表面",
    },
}


def check_template(root: Path, relative: str, spec: dict) -> list[str]:
    errors: list[str] = []
    path = root / Path(relative)
    if not path.is_file():
        return [f"missing: {relative}"]
    expected = EXPECTED[relative]
    if path.stat().st_size != expected["bytes"] or sha256(path) != expected["sha256"]:
        errors.append(f"official template hash/size changed: {relative}")

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if tuple(workbook.sheetnames) != spec["sheets"]:
            errors.append(f"sheet names: {relative} expected={spec['sheets']} actual={tuple(workbook.sheetnames)}")
        for worksheet in workbook.worksheets:
            if (worksheet.max_row, worksheet.max_column) != (5, 6):
                errors.append(f"dimensions: {relative}/{worksheet.title} expected=(5, 6) actual=({worksheet.max_row}, {worksheet.max_column})")
            rows = list(worksheet.iter_rows(min_row=1, max_row=5, min_col=1, max_col=6, values_only=True))
            if not rows or rows[0][0] != "时间\\到药材中心的距离":
                errors.append(f"header label: {relative}/{worksheet.title}")
            if len(rows) >= 1 and (rows[0][1], rows[0][2], rows[0][3], rows[0][5]) != (0, 0.1, 0.2, spec["last_header"]):
                errors.append(f"distance header: {relative}/{worksheet.title}")
            if len(rows) == 5:
                time_values = tuple(row[0] for row in rows[1:])
                if time_values != spec["time_examples"]:
                    errors.append(f"time skeleton: {relative}/{worksheet.title} expected={spec['time_examples']} actual={time_values}")
                for row_number, row in enumerate(rows[1:], start=2):
                    if any(value is not None for value in row[1:]):
                        errors.append(f"template data area is not empty: {relative}/{worksheet.title} row={row_number}")
    finally:
        workbook.close()
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    root = args.root.resolve()
    errors: list[str] = []
    for relative, spec in TEMPLATES.items():
        found = check_template(root, relative, spec)
        if found:
            errors.extend(found)
            for error in found:
                print(f"[FAIL] {error}")
        else:
            print(f"[PASS] {relative}")
    if errors:
        print(f"Template validation failed: {len(errors)} issue(s)")
        return 1
    print(f"Template validation passed: {len(TEMPLATES)} official template(s)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
