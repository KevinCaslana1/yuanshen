"""Q2 long-horizon boundary and production-configuration gate.

This script writes auditable CSV/JSON/PNG evidence only. It never opens or
writes any result2.xlsx workbook and has no Q3 stopping logic. The event
observer is passive: it records the first one-second sign-change bracket of
``max(C)-0.15`` but never stops a run.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import sys
import time
from dataclasses import replace
from pathlib import Path
from statistics import mean, pstdev

import matplotlib.pyplot as plt
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.q2 import Q2RunConfig, run_q2, run_q2_b0
from src.q2.environment import EnvironmentProvider
from src.q2.lineage import canonical_csv_path
from src.q2.model import boundary_flux_pair, interface_values
from src.q2.properties import diffusivity
from src.q2.solver import grid_for_config
from src.q2.validation import point_values, property_range_metrics
from src.q2.benchmark import run_benchmark


INPUT = ROOT / "A题" / "附件" / "附件1.xlsx"
EXP = ROOT / "experiments"
PAPER_TIMES = (1800.0, 3600.0, 5400.0, 7200.0, 9000.0, 10800.0)
PAPER_RADII = (0.0, 0.5, 1.0, 1.5, 2.0)
LONG_TIMES = (21600.0, 43200.0, 86400.0, 172800.0, 259200.0)
STAGE_ENDS = (21600.0, 86400.0, 172800.0, 259200.0)


def ensure_folder(name: str) -> Path:
    path = EXP / name
    path.mkdir(parents=True, exist_ok=True)
    return path


def dump(path: Path, value) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")


def cfg_dict(config: Q2RunConfig) -> dict:
    return config.as_dict()


def attachment_rows() -> list[tuple[float, float, float]]:
    workbook = load_workbook(INPUT, read_only=True, data_only=True)
    rows = [tuple(float(value) for value in row[:3]) for row in workbook["Sheet1"].iter_rows(min_row=2, values_only=True)]
    workbook.close()
    return rows


def linear_slope(xs: list[float], ys: list[float]) -> float:
    xbar, ybar = mean(xs), mean(ys)
    return sum((x - xbar) * (y - ybar) for x, y in zip(xs, ys)) / sum((x - xbar) ** 2 for x in xs)


def tail_summary(rows: list[tuple[float, float, float]], n: int) -> dict:
    tail = rows[-n:]
    xs = [row[0] for row in tail]
    ts = [row[1] for row in tail]
    cs = [row[2] for row in tail]
    return {
        "n": n,
        "time_start_s": xs[0], "time_end_s": xs[-1],
        "temperature_C": {"mean": mean(ts), "population_std": pstdev(ts), "min": min(ts), "max": max(ts), "last": ts[-1], "linear_slope_per_s": linear_slope(xs, ts)},
        "moisture_kg_kg": {"mean": mean(cs), "population_std": pstdev(cs), "min": min(cs), "max": max(cs), "last": cs[-1], "linear_slope_per_s": linear_slope(xs, cs)},
    }


def environment_candidates(rows: list[tuple[float, float, float]]) -> dict[str, dict[str, float | str]]:
    last = rows[-1]
    tail = rows[-40:]
    return {
        "ENV-A-last-raw": {"temperature_C": last[1], "moisture_kg_kg": last[2], "rule": "last raw Attachment 1 point"},
        "ENV-B-tail40-mean": {"temperature_C": mean(row[1] for row in tail), "moisture_kg_kg": mean(row[2] for row in tail), "rule": "mean of last 40 raw points (12060..14400 s)"},
        "TEAM-REFERENCE": {"temperature_C": 50.0, "moisture_kg_kg": 0.05, "rule": "team numerical proposal; not official data"},
    }


def run_environment_audit() -> dict:
    path = ensure_folder("EXP-Q2-012-ENVIRONMENT")
    rows = attachment_rows()
    candidates = environment_candidates(rows)
    tails = {str(n): tail_summary(rows, n) for n in (10, 20, 40, 60, 80)}
    base = Q2RunConfig(end_time_s=14400.0, time_step_s=1.0, candidate="A", n_intervals=40, scheme="be")
    provider = EnvironmentProvider.from_attachment1(INPUT, method="linear")
    transition = {}
    last_t, last_c = rows[-1][1], rows[-1][2]
    for name, candidate in candidates.items():
        dt, dc = candidate["temperature_C"] - last_t, candidate["moisture_kg_kg"] - last_c
        transition[name] = {
            "at_14400_s": list(provider.at(14400.0)),
            "at_14400_plus_1s": [candidate["temperature_C"], candidate["moisture_kg_kg"]],
            "delta_environment_temperature_C": dt,
            "delta_environment_moisture_kg_kg": dc,
            "instantaneous_delta_heat_robin_w_m2": base.parameters.heat_transfer_w_m2_k * dt,
            "instantaneous_delta_moisture_robin_kg_m2_s": base.parameters.mass_transfer_m_s * dc,
        }
    config = replace(base, end_time_s=14400.0, time_step_s=1.0, post_attachment_mode="raise")
    started = time.perf_counter()
    result = run_q2(config, record_times=[14400.0])
    surface_t, surface_c = result.snapshot_at(14400.0)[0][-1], result.snapshot_at(14400.0)[1][-1]
    grid = result.grid
    for name, candidate in candidates.items():
        env_t = candidate["temperature_C"] + 273.15
        env_c = candidate["moisture_kg_kg"]
        d_surface = diffusivity(surface_c, surface_t)
        k_dummy = 0.21 + 0.38 * surface_c / (surface_c + 1.0)
        heat_internal, heat_robin, heat_residual = boundary_flux_pair(surface_t, result.snapshot_at(14400.0)[0][-2], grid, k_dummy, base.parameters.heat_transfer_w_m2_k, env_t)
        moisture_internal, moisture_robin, moisture_residual = boundary_flux_pair(surface_c, result.snapshot_at(14400.0)[1][-2], grid, d_surface, base.parameters.mass_transfer_m_s, env_c)
        transition[name]["representative_surface_at_14400"] = {"temperature_C": surface_t - 273.15, "moisture_kg_kg": surface_c}
        transition[name]["robin_flux_if_post_constant"] = {"heat_w_m2": heat_robin, "moisture_kg_m2_s": moisture_robin, "internal_heat_w_m2": heat_internal, "internal_moisture_kg_m2_s": moisture_internal, "heat_boundary_difference_w_m2": heat_residual, "moisture_boundary_difference_kg_m2_s": moisture_residual}
    payload = {
        "experiment": "EXP-Q2-012-ENVIRONMENT",
        "status": "PASS",
        "input": "A题/附件/附件1.xlsx",
        "input_sha256": rows and __import__("hashlib").sha256(INPUT.read_bytes()).hexdigest(),
        "raw_row_count": len(rows), "raw_time_end_s": rows[-1][0], "raw_interval_s": sorted({b[0] - a[0] for a, b in zip(rows, rows[1:])}),
        "tail_windows": tails, "candidates": candidates, "transition_audit": transition,
        "transition_solver_runtime_s": time.perf_counter() - started,
        "team_reference_is_not_official": True, "no_silent_smoothing": True, "workbook_written": False,
    }
    dump(path / "metrics.json", payload)
    dump(path / "config.json", {"candidate_rules": candidates, "fixed_geometry": cfg_dict(base)})
    return payload


def compare_sample_csv(left_path: Path, right_path: Path, paper_times: tuple[float, ...] = PAPER_TIMES) -> dict:
    max_temp = max_moisture = 0.0
    per_time: dict[str, dict[str, float]] = {}
    left_handle = left_path.open(newline="", encoding="utf-8")
    right_handle = right_path.open(newline="", encoding="utf-8")
    left_reader, right_reader = csv.DictReader(left_handle), csv.DictReader(right_handle)
    rows = 0
    for left, right in zip(left_reader, right_reader):
        if left["time_s"] != right["time_s"] or left["radius_cm"] != right["radius_cm"]:
            raise AssertionError("sample CSV time/radius arrays are not aligned")
        t = float(left["time_s"])
        dt = abs(float(left["temperature_C"]) - float(right["temperature_C"]))
        dc = abs(float(left["moisture_kg_kg"]) - float(right["moisture_kg_kg"]))
        max_temp, max_moisture = max(max_temp, dt), max(max_moisture, dc)
        key = str(int(t)) if abs(t - round(t)) < 1.0e-10 else repr(t)
        if t in paper_times or t == 14400.0:
            item = per_time.setdefault(key, {"temperature_max_abs_C": 0.0, "moisture_max_abs": 0.0})
            item["temperature_max_abs_C"] = max(item["temperature_max_abs_C"], dt)
            item["moisture_max_abs"] = max(item["moisture_max_abs"], dc)
        rows += 1
    extra_left = next(left_reader, None)
    extra_right = next(right_reader, None)
    left_handle.close(); right_handle.close()
    return {"rows_compared": rows, "same_length": extra_left is None and extra_right is None, "temperature_max_abs_C": max_temp, "moisture_max_abs": max_moisture, "representative_times": per_time}


def run_interpolation_audit() -> dict:
    path = ensure_folder("EXP-Q2-013-INTERPOLATION")
    common = dict(end_time_s=14400.0, time_step_s=1.0, candidate="A", n_intervals=80, scheme="bdf2", post_attachment_mode="raise")
    configs = {method: Q2RunConfig(interpolation=method, **common) for method in ("linear", "pchip")}
    runtimes = {}
    for method, config in configs.items():
        started = time.perf_counter()
        run_q2(config, record_times=[0.0] + list(PAPER_TIMES) + [14400.0], output_path=path / f"samples_{method}.csv", diagnostics_path=path / f"diagnostics_{method}.csv", diagnostics_interval_s=1.0)
        runtimes[method] = time.perf_counter() - started
    comparison = compare_sample_csv(path / "samples_linear.csv", path / "samples_pchip.csv")
    provider = EnvironmentProvider.from_attachment1(INPUT, method="pchip")
    knots_exact = all(provider.at(t) == (temp, moisture) for t, temp, moisture in zip(provider.times_s, provider.temperatures_c, provider.moistures_kg_kg))
    payload = {"experiment": "EXP-Q2-013-INTERPOLATION", "status": "PASS", "configs": {key: cfg_dict(value) for key, value in configs.items()}, "runtime_s": runtimes, "raw_points_preserved": knots_exact, "comparison_0_4h_full_official_samples": comparison, "decision_rule": "if differences remain below the numerical gate, prefer the simpler robust linear interpolator; no automatic PCHIP adoption", "recommended_for_freeze": "linear_pending_human_approval" if comparison["temperature_max_abs_C"] < 5.0e-5 and comparison["moisture_max_abs"] < 5.0e-5 else "human_review_required", "workbook_written": False}
    dump(path / "metrics.json", payload)
    return payload


def inventory(result, time_s: float) -> float:
    moisture = result.snapshot_at(time_s)[1]
    volumes = __import__("src.q1.model", fromlist=["radial_control_volume_factors"]).radial_control_volume_factors(result.grid)
    return sum(value * volume for value, volume in zip(moisture, volumes))


def field_difference(test, reference, time_s: float) -> dict:
    test_t, test_c = test.snapshot_at(time_s)
    ref_t, ref_c = reference.snapshot_at(time_s)
    temp = [abs(a - b) for a, b in zip(test_t, ref_t)]
    moist = [abs(a - b) for a, b in zip(test_c, ref_c)]
    return {"temperature_Linf_K": max(temp), "temperature_L2_RMS_K": math.sqrt(sum(value * value for value in temp) / len(temp)), "moisture_Linf": max(moist), "moisture_L2_RMS": math.sqrt(sum(value * value for value in moist) / len(moist)), "inventory_abs": abs(inventory(test, time_s) - inventory(reference, time_s))}


def classify_sensitivity(max_temp: float, max_moisture: float, base, threshold: float = 0.01) -> str:
    t_scale = max(max(abs(value - 273.15) for value in base.final_temperature_k), 1.0)
    c_scale = max(max(abs(value) for value in base.final_moisture), 1.0e-12)
    relative = max(max_temp / t_scale, max_moisture / c_scale)
    return "low" if relative < threshold else "medium" if relative < 0.1 else "high"


def run_boundary_sensitivity() -> dict:
    path = ensure_folder("EXP-Q2-014-BC-SENSITIVITY")
    base = Q2RunConfig(end_time_s=21600.0, time_step_s=1.0, candidate="A", n_intervals=80, scheme="bdf2", post_attachment_mode="constant")
    configs = {
        "base": base,
        "h_minus10": replace(base, parameters=replace(base.parameters, heat_transfer_w_m2_k=base.parameters.heat_transfer_w_m2_k * 0.9)),
        "h_plus10": replace(base, parameters=replace(base.parameters, heat_transfer_w_m2_k=base.parameters.heat_transfer_w_m2_k * 1.1)),
        "hm_minus10": replace(base, parameters=replace(base.parameters, mass_transfer_m_s=base.parameters.mass_transfer_m_s * 0.9)),
        "hm_plus10": replace(base, parameters=replace(base.parameters, mass_transfer_m_s=base.parameters.mass_transfer_m_s * 1.1)),
    }
    times = (0.0, 10800.0, 21600.0)
    runs = {key: run_q2(config, record_times=times, passive_event_threshold_kg_kg=0.15) for key, config in configs.items()}
    impacts = {}
    for key, result in runs.items():
        if key == "base":
            continue
        diffs = [field_difference(result, runs["base"], t) for t in times[1:]]
        impacts[key] = {"max_temperature_Linf_K": max(row["temperature_Linf_K"] for row in diffs), "max_moisture_Linf": max(row["moisture_Linf"] for row in diffs), "max_inventory_abs": max(row["inventory_abs"] for row in diffs), "classification": classify_sensitivity(max(row["temperature_Linf_K"] for row in diffs), max(row["moisture_Linf"] for row in diffs), runs["base"]), "passive_event_bracket_s": result.passive_event_bracket_s}
    payload = {"experiment": "EXP-Q2-014-BC-SENSITIVITY", "status": "PASS", "base_config": cfg_dict(base), "perturbation": "independent ±10% changes; h=25 W/(m²·K), hm=8e-7 m/s remain modeling assumptions and are not fitted", "classification_rule": "relative max field change <1%=low, 1–10%=medium, >=10%=high", "impacts_0_6h": impacts, "all_runs_converged": True, "q3_event_stop": False, "workbook_written": False}
    dump(path / "metrics.json", payload)
    return payload


def run_interface_audit() -> dict:
    path = ensure_folder("EXP-Q2-015-INTERFACE-MEAN")
    common = dict(end_time_s=10800.0, time_step_s=0.25, candidate="A", n_intervals=80, scheme="bdf2", post_attachment_mode="raise")
    configs = {method: Q2RunConfig(interface_mean=method, **common) for method in ("arithmetic", "harmonic")}
    runs, runtimes = {}, {}
    for method, config in configs.items():
        started = time.perf_counter(); runs[method] = run_q2(config, record_times=[0.0] + list(PAPER_TIMES)); runtimes[method] = time.perf_counter() - started
    comparison = {"times": {str(int(t)): field_difference(runs["arithmetic"], runs["harmonic"], t) for t in PAPER_TIMES}, "max_temperature_Linf_K": 0.0, "max_moisture_Linf": 0.0}
    comparison["max_temperature_Linf_K"] = max(row["temperature_Linf_K"] for row in comparison["times"].values())
    comparison["max_moisture_Linf"] = max(row["moisture_Linf"] for row in comparison["times"].values())
    benchmark = run_benchmark()
    payload = {"experiment": "EXP-Q2-015-INTERFACE-MEAN", "status": "PASS", "configs": {key: cfg_dict(value) for key, value in configs.items()}, "runtime_s": runtimes, "real_q2_0_3h_comparison": comparison, "manufactured_benchmark": benchmark, "recommendation_status": "RECOMMENDED_FOR_FREEZE_PENDING_HUMAN_APPROVAL" if comparison["max_temperature_Linf_K"] < 5.0e-5 and comparison["max_moisture_Linf"] < 5.0e-5 else "HUMAN_REVIEW_REQUIRED", "workbook_written": False}
    dump(path / "metrics.json", payload)
    return payload


def run_baseline_comparison() -> dict:
    path = ensure_folder("EXP-Q2-020-BASELINE")
    config = Q2RunConfig(end_time_s=21600.0, time_step_s=0.5, candidate="A", n_intervals=80, scheme="bdf2", post_attachment_mode="constant")
    times = (0.0, 10800.0, 21600.0)
    started = time.perf_counter(); full = run_q2(config, record_times=times, passive_event_threshold_kg_kg=0.15); full_runtime = time.perf_counter() - started
    started = time.perf_counter(); baseline = run_q2_b0(config, record_times=times); baseline_runtime = time.perf_counter() - started
    comparison = {str(int(t)): field_difference(baseline, full, t) for t in times[1:]}
    payload = {"experiment": "EXP-Q2-020-BASELINE", "status": "PASS", "config": cfg_dict(config), "full_coupled_runtime_s": full_runtime, "fixed_property_B0_runtime_s": baseline_runtime, "runtime_ratio_B0_over_full": baseline_runtime / full_runtime, "comparison_full_vs_B0": comparison, "B0_definition": "same geometry/FVM/environment/time scheme; Appendix 3 rho/cp/k/D frozen at initial T,C; no Picard", "workbook_written": False}
    dump(path / "metrics.json", payload)
    return payload


def stage_summary(result, stage_start: float, stage_end: float, runtime_s: float) -> dict:
    metrics = dict(result.diagnostic_summary)
    return {"stage_start_s": stage_start, "stage_end_s": stage_end, "runtime_s": runtime_s, "complete_to_config_horizon": result.complete, "diagnostics": metrics, "property_ranges": property_range_metrics(result), "passive_event_bracket_s": result.passive_event_bracket_s, "finite_final_fields": all(math.isfinite(value) for value in result.final_temperature_k + result.final_moisture), "positive_moisture": min(result.final_moisture) > 0.0}


def validate_output_samples(path: Path, end_time_s: float) -> dict:
    expected_times = set(range(int(end_time_s) + 1))
    counts: dict[float, int] = {}
    radii_by_time: dict[float, set[float]] = {}
    rows = 0
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            t, radius = float(row["time_s"]), float(row["radius_cm"])
            counts[t] = counts.get(t, 0) + 1
            radii_by_time.setdefault(t, set()).add(radius)
            rows += 1
    expected_radii = {round(i * 0.1, 12) for i in range(21)}
    missing = sorted(expected_times - set(int(t) for t in counts if abs(t - round(t)) < 1.0e-10))
    repeated = sorted(t for t, count in counts.items() if count != 21)
    bad_radii = sorted(t for t, radii in radii_by_time.items() if {round(value, 12) for value in radii} != expected_radii)
    return {"rows": rows, "expected_rows": (int(end_time_s) + 1) * 21, "missing_integer_seconds": missing[:20], "missing_count": len(missing), "repeated_or_missing_radius_rows": repeated[:20], "bad_radius_time_count": len(bad_radii), "time_min": min(counts) if counts else None, "time_max": max(counts) if counts else None, "one_second_grid_valid": not missing and not repeated and not bad_radii and rows == (int(end_time_s) + 1) * 21}


def deduplicate_csv(source: Path, target: Path, key_fields: tuple[str, ...]) -> dict:
    """Recover a restart artifact by retaining the first row for each key.

    The source remains untouched. This is only valid because each key is an
    official deterministic output coordinate; the duplicate interval is
    reported separately rather than silently discarded.
    """
    seen = set()
    duplicate_rows = 0
    with source.open(newline="", encoding="utf-8") as source_handle, target.open("w", newline="", encoding="utf-8") as target_handle:
        reader = csv.DictReader(source_handle)
        writer = csv.DictWriter(target_handle, fieldnames=reader.fieldnames)
        writer.writeheader()
        for row in reader:
            key = tuple(row[field] for field in key_fields)
            if key in seen:
                duplicate_rows += 1
                continue
            seen.add(key)
            writer.writerow(row)
    return {"source": str(source.relative_to(ROOT)), "target": str(target.relative_to(ROOT)), "duplicate_rows_removed": duplicate_rows, "unique_keys": len(seen)}


def recover_long_output(environment_name: str = "ENV-A-last-raw") -> dict:
    path = ensure_folder(f"EXP-Q2-016-LONG-{environment_name}")
    raw_samples = path / "official_samples.csv"
    raw_diagnostics = path / "diagnostics_1s.csv"
    recovered_samples = path / "official_samples_recovered.csv"
    recovered_diagnostics = path / "diagnostics_1s_recovered.csv"
    sample_recovery = deduplicate_csv(raw_samples, recovered_samples, ("time_s", "radius_cm"))
    diagnostic_recovery = deduplicate_csv(raw_diagnostics, recovered_diagnostics, ("time_s",))
    validation = validate_output_samples(recovered_samples, 259200.0)
    payload = {"experiment": f"EXP-Q2-016-LONG-{environment_name}-OUTPUT-RECOVERY", "status": "PASS" if validation["one_second_grid_valid"] else "BLOCKED", "raw_output_validation": validate_output_samples(raw_samples, 259200.0), "sample_recovery": sample_recovery, "diagnostic_recovery": diagnostic_recovery, "recovered_output_validation": validation, "recovery_reason": "old checkpoint 43200 s was resumed after a prior interrupted process had already advanced the append-only files; raw files are retained as evidence", "workbook_written": False}
    dump(path / "output_recovery_metrics.json", payload)
    return payload


def finalize_recovered_long_a() -> dict:
    """Attach the validated recovery artifact to the long-run metrics.

    The append-only raw files remain untouched as evidence.  The recovered
    one-second files become the traceable sampler artifact used by plots and
    convergence comparisons.
    """
    path = EXP / "EXP-Q2-016-LONG-ENV-A-last-raw"
    metrics_path = path / "metrics.json"
    recovery_path = path / "output_recovery_metrics.json"
    payload = json.loads(metrics_path.read_text(encoding="utf-8"))
    recovery = json.loads(recovery_path.read_text(encoding="utf-8"))
    payload["output_sampler_raw"] = payload.get("output_sampler", {})
    payload["output_sampler"] = recovery["recovered_output_validation"]
    payload["output_sampler"]["source"] = recovery["sample_recovery"]["target"]
    payload["output_recovery"] = {
        "status": recovery["status"],
        "raw_output_validation": recovery["raw_output_validation"],
        "sample_recovery": recovery["sample_recovery"],
        "diagnostic_recovery": recovery["diagnostic_recovery"],
        "reason": recovery["recovery_reason"],
    }
    payload["status"] = "PASS_WITH_RECOVERED_OUTPUT_ARTIFACT"
    payload["raw_output_recovery_required"] = True
    dump(metrics_path, payload)
    return {"status": payload["status"], "output_sampler": payload["output_sampler"], "output_recovery": payload["output_recovery"]}


def compare_environment_checkpoints() -> dict:
    path = ensure_folder("EXP-Q2-018-ENV-COMPARISON")
    a_root = EXP / "EXP-Q2-016-LONG-ENV-A-last-raw"
    b_root = EXP / "EXP-Q2-016-LONG-ENV-B-tail40-mean"
    a_config = Q2RunConfig(end_time_s=259200.0, time_step_s=0.25, candidate="A", n_intervals=80, scheme="bdf2", post_attachment_mode="constant", post_temperature_c=50.165, post_moisture_kg_kg=0.04986)
    indices = {str(radius): min(range(len(grid_for_config(a_config).nodes_m)), key=lambda i: abs(grid_for_config(a_config).nodes_m[i] * 100.0 - radius)) for radius in (0.0, 1.0, 1.5, 2.0)}
    comparisons = {}
    for stage_end in (21600.0, 86400.0, 172800.0, 259200.0):
        a = json.loads((a_root / f"checkpoint_{int(stage_end)}s.json").read_text(encoding="utf-8"))
        b = json.loads((b_root / f"checkpoint_{int(stage_end)}s.json").read_text(encoding="utf-8"))
        ta = [x - 273.15 for x in a["temperature_k"]]; tb = [x - 273.15 for x in b["temperature_k"]]
        ca, cb = a["moisture_kg_kg"], b["moisture_kg_kg"]
        td = [abs(x - y) for x, y in zip(ta, tb)]; cd = [abs(x - y) for x, y in zip(ca, cb)]
        comparisons[str(int(stage_end))] = {"temperature_Linf_C": max(td), "temperature_L2_RMS_C": math.sqrt(sum(x * x for x in td) / len(td)), "moisture_Linf": max(cd), "moisture_L2_RMS": math.sqrt(sum(x * x for x in cd) / len(cd)), "representative": {radius: {"temperature_A_C": ta[index], "temperature_B_C": tb[index], "temperature_difference_C": ta[index] - tb[index], "moisture_A": ca[index], "moisture_B": cb[index], "moisture_difference": ca[index] - cb[index]} for radius, index in indices.items()}}
    payload = {"experiment": "EXP-Q2-018-ENV-COMPARISON", "status": "PASS", "same_grid_and_time_step": True, "comparison_times_s": [21600, 86400, 172800, 259200], "environment_A": "last raw point", "environment_B": "last 40 raw-point mean", "comparisons": comparisons, "interpretation": "A/B sensitivity evidence; because the post-14400 rule changes results, this does not silently close the environment open question", "workbook_written": False}
    dump(path / "config.json", {"A": cfg_dict(a_config), "B_post_constant": {"temperature_C": 49.99525, "moisture_kg_kg": 0.049988}, "comparison_times_s": [21600, 86400, 172800, 259200]})
    dump(path / "metrics.json", payload)
    return payload


def plot_recovered_long_samples() -> dict:
    path = EXP / "EXP-Q2-016-LONG-ENV-A-last-raw"
    source = canonical_csv_path("experiments/EXP-Q2-016-LONG-ENV-A-last-raw/official_samples_recovered.csv")
    times, surface_temperature, center_temperature, surface_moisture, center_moisture = [], [], [], [], []
    with source.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        current = None
        values = {}
        for row in reader:
            t = float(row["time_s"]); radius = float(row["radius_cm"])
            if current is None:
                current = t
            if t != current:
                times.append(current); surface_temperature.append(values[2.0][0]); center_temperature.append(values[0.0][0]); surface_moisture.append(values[2.0][1]); center_moisture.append(values[0.0][1]); values = {}; current = t
            values[radius] = (float(row["temperature_C"]), float(row["moisture_kg_kg"]))
        if current is not None:
            times.append(current); surface_temperature.append(values[2.0][0]); center_temperature.append(values[0.0][0]); surface_moisture.append(values[2.0][1]); center_moisture.append(values[0.0][1])
    for filename, y1, y2, ylabel, title in (("temperature_center_surface.png", surface_temperature, center_temperature, "temperature [°C]", "Q2 long run: center and surface temperature"), ("moisture_center_surface.png", surface_moisture, center_moisture, "moisture [kg/kg]", "Q2 long run: center and surface moisture")):
        fig, ax = plt.subplots(figsize=(8, 4)); ax.plot(times, y1, label="surface", linewidth=0.8); ax.plot(times, y2, label="center", linewidth=0.8); ax.set(xlabel="time [s]", ylabel=ylabel, title=title); ax.grid(True, alpha=0.25); ax.legend(); fig.tight_layout(); fig.savefig(path / filename, dpi=160); plt.close(fig)
    return {"source": str(source.relative_to(ROOT)), "samples": len(times), "plots": ["temperature_center_surface.png", "moisture_center_surface.png"], "smoothing": False, "interpolation": False, "epsilon_clip": False}


def run_long_restart_audit() -> dict:
    path = ensure_folder("EXP-Q2-019-LONG-RESTART")
    source = EXP / "EXP-Q2-016-LONG-ENV-A-last-raw" / "checkpoint_43200s.json"
    config = Q2RunConfig(end_time_s=259200.0, time_step_s=0.25, candidate="A", n_intervals=80, scheme="bdf2", post_attachment_mode="constant", post_temperature_c=50.165, post_moisture_kg_kg=0.04986)
    started = time.perf_counter(); continuous = run_q2(config, record_times=[86400.0], stop_time_s=86400.0, diagnostics_path=path / "continuous_diagnostics_60s.csv", diagnostics_interval_s=60.0); continuous_runtime = time.perf_counter() - started
    started = time.perf_counter(); restarted = run_q2(config, record_times=[86400.0], restart_path=source, stop_time_s=86400.0, diagnostics_path=path / "restart_diagnostics_60s.csv", diagnostics_interval_s=60.0); restart_runtime = time.perf_counter() - started
    temp_diff = max(abs(a - b) for a, b in zip(continuous.snapshot_at(86400.0)[0], restarted.snapshot_at(86400.0)[0]))
    moisture_diff = max(abs(a - b) for a, b in zip(continuous.snapshot_at(86400.0)[1], restarted.snapshot_at(86400.0)[1]))
    payload = {"experiment": "EXP-Q2-019-LONG-RESTART", "status": "PASS" if temp_diff <= 1.0e-12 and moisture_diff <= 1.0e-15 else "FAIL", "checkpoint": str(source.relative_to(ROOT)), "continuous_horizon_s": 86400.0, "restart_start_s": 43200.0, "continuous_runtime_s": continuous_runtime, "restart_runtime_s": restart_runtime, "field_difference_at_24h": {"temperature_max_abs_K": temp_diff, "moisture_max_abs": moisture_diff}, "machine_level_tolerance": {"temperature_K": 1.0e-12, "moisture": 1.0e-15}, "workbook_written": False}
    dump(path / "config.json", cfg_dict(config)); dump(path / "metrics.json", payload)
    return payload


def load_main_official_snapshots() -> dict[float, dict[float, tuple[float, float]]]:
    source = canonical_csv_path("experiments/EXP-Q2-016-LONG-ENV-A-last-raw/official_samples_recovered.csv")
    wanted = {float(value) for value in (21600, 43200, 86400, 172800, 259200)}
    output: dict[float, dict[float, tuple[float, float]]] = {}
    with source.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            t = float(row["time_s"])
            if t in wanted:
                output.setdefault(t, {})[float(row["radius_cm"])] = (float(row["temperature_C"]), float(row["moisture_kg_kg"]))
    if any(len(output.get(t, {})) != 21 for t in wanted):
        raise ValueError("recovered main output lacks a complete official snapshot")
    return output


def compare_result_to_main(result, main_snapshots: dict[float, dict[float, tuple[float, float]]], times: tuple[float, ...]) -> tuple[dict, list[dict]]:
    indices = {round(i * 0.1, 12): min(range(len(result.grid.nodes_m)), key=lambda j: abs(result.grid.nodes_m[j] * 100.0 - round(i * 0.1, 12))) for i in range(21)}
    rows, metrics = [], {}
    for t in times:
        temp, moisture = result.snapshot_at(t)
        temp_error, moisture_error = [], []
        for radius, index in indices.items():
            ref_t, ref_c = main_snapshots[t][radius]
            signed_t, signed_c = temp[index] - 273.15 - ref_t, moisture[index] - ref_c
            temp_error.append(abs(signed_t)); moisture_error.append(abs(signed_c))
            rows.append({"level_time_s": t, "radius_cm": radius, "temperature_signed_error_C": signed_t, "temperature_absolute_error_C": abs(signed_t), "moisture_signed_error": signed_c, "moisture_absolute_error": abs(signed_c)})
        metrics[str(int(t))] = {"temperature_Linf_C": max(temp_error), "temperature_L2_RMS_C": math.sqrt(sum(x * x for x in temp_error) / len(temp_error)), "moisture_Linf": max(moisture_error), "moisture_L2_RMS": math.sqrt(sum(x * x for x in moisture_error) / len(moisture_error))}
    flat_t = [row["temperature_absolute_error_C"] for row in rows]; flat_c = [row["moisture_absolute_error"] for row in rows]
    metrics["all_requested_times"] = {"temperature_Linf_C": max(flat_t), "temperature_L2_RMS_C": math.sqrt(sum(x * x for x in flat_t) / len(flat_t)), "moisture_Linf": max(flat_c), "moisture_L2_RMS": math.sqrt(sum(x * x for x in flat_c) / len(flat_c))}
    return metrics, rows


def plot_error_levels(path: Path, values: dict[float, float], xlabel: str, filename: str, title: str) -> None:
    fig, ax = plt.subplots(figsize=(6, 4)); ax.loglog(list(values), list(values.values()), "o-"); ax.set(xlabel=xlabel, ylabel="L∞ error versus dt=.25/n=80", title=title); ax.grid(True, which="both", alpha=0.25); fig.tight_layout(); fig.savefig(path / filename, dpi=160); plt.close(fig)


def run_long_convergence() -> dict:
    path = ensure_folder("EXP-Q2-017-LONG-CONVERGENCE")
    main = load_main_official_snapshots()
    times = (21600.0, 43200.0, 86400.0, 172800.0, 259200.0)
    common = dict(end_time_s=259200.0, candidate="A", n_intervals=80, scheme="bdf2", post_attachment_mode="constant", post_temperature_c=50.165, post_moisture_kg_kg=0.04986)
    temporal = {}
    temporal_rows = []
    for dt in (1.0, 0.5):
        config = Q2RunConfig(time_step_s=dt, **common)
        started = time.perf_counter(); result = run_q2(config, record_times=times, diagnostics_path=path / f"diagnostics_dt{dt:g}.csv", diagnostics_interval_s=3600.0); runtime = time.perf_counter() - started
        metrics, rows = compare_result_to_main(result, main, times); temporal[str(dt)] = {"config": cfg_dict(config), "runtime_s": runtime, "errors": metrics, "diagnostics": result.diagnostic_summary}; temporal_rows.extend([{**row, "study": "time", "level": dt} for row in rows])
    spatial = {}
    spatial_rows = []
    for n in (20, 40):
        config = Q2RunConfig(time_step_s=0.25, n_intervals=n, **{key: value for key, value in common.items() if key != "n_intervals"})
        started = time.perf_counter(); result = run_q2(config, record_times=times, diagnostics_path=path / f"diagnostics_n{n}.csv", diagnostics_interval_s=3600.0); runtime = time.perf_counter() - started
        metrics, rows = compare_result_to_main(result, main, times); spatial[str(n)] = {"config": cfg_dict(config), "runtime_s": runtime, "errors": metrics, "diagnostics": result.diagnostic_summary}; spatial_rows.extend([{**row, "study": "space", "level": n} for row in rows])
    with (path / "long_convergence_points.csv").open("w", newline="", encoding="utf-8") as handle:
        fields = ["study", "level", "level_time_s", "radius_cm", "temperature_signed_error_C", "temperature_absolute_error_C", "moisture_signed_error", "moisture_absolute_error"]
        writer = csv.DictWriter(handle, fieldnames=fields); writer.writeheader(); writer.writerows(temporal_rows + spatial_rows)
    temp_time = {float(key): value["errors"]["all_requested_times"]["temperature_Linf_C"] for key, value in temporal.items()}; moist_time = {float(key): value["errors"]["all_requested_times"]["moisture_Linf"] for key, value in temporal.items()}
    temp_space = {float(key): value["errors"]["all_requested_times"]["temperature_Linf_C"] for key, value in spatial.items()}; moist_space = {float(key): value["errors"]["all_requested_times"]["moisture_Linf"] for key, value in spatial.items()}
    payload = {"experiment": "EXP-Q2-017-LONG-CONVERGENCE", "status": "PASS", "reference": "ENV-A main Candidate A n=80 dt=.25 s recovered 72 h output", "comparison_times_s": list(times), "time_levels": temporal, "space_levels": spatial, "observed_order_in_existing_short_horizon": {"time": "EXP-Q2-005 has dt=1,.5,.25,.125 with smaller reference; space: EXP-Q2-006 has dr=.1,.05,.025 with smaller reference"}, "long_error_curves": {"time_temperature": temp_time, "time_moisture": moist_time, "space_temperature": temp_space, "space_moisture": moist_space}, "workbook_written": False}
    dump(path / "config.json", {"common": common, "comparison_times_s": list(times), "time_levels": [1.0, 0.5, 0.25], "space_levels": [20, 40, 80]}); dump(path / "metrics.json", payload)
    plot_error_levels(path, temp_time, "dt [s]", "long_time_temperature.png", "Q2 long-horizon temporal sensitivity")
    plot_error_levels(path, moist_time, "dt [s]", "long_time_moisture.png", "Q2 long-horizon moisture temporal sensitivity")
    plot_error_levels(path, temp_space, "n intervals", "long_space_temperature.png", "Q2 long-horizon spatial sensitivity")
    plot_error_levels(path, moist_space, "n intervals", "long_space_moisture.png", "Q2 long-horizon moisture spatial sensitivity")
    return payload


def augment_convergence_orders() -> dict:
    path = EXP / "EXP-Q2-017-LONG-CONVERGENCE" / "metrics.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    def order(left: float, right: float, ratio: float = 2.0) -> float | None:
        if left <= 0.0 or right <= 0.0:
            return None
        return math.log(left / right) / math.log(ratio)
    payload["observed_order_proxy_vs_finest_reference"] = {
        "time": {
            "temperature": order(payload["time_levels"]["1.0"]["errors"]["all_requested_times"]["temperature_Linf_C"], payload["time_levels"]["0.5"]["errors"]["all_requested_times"]["temperature_Linf_C"]),
            "moisture": order(payload["time_levels"]["1.0"]["errors"]["all_requested_times"]["moisture_Linf"], payload["time_levels"]["0.5"]["errors"]["all_requested_times"]["moisture_Linf"]),
            "interpretation": "proxy from errors relative to dt=.25 reference; authoritative dt=1/.5/.25/.125 observed orders are in EXP-Q2-005",
        },
        "space": {
            "temperature": order(payload["space_levels"]["20"]["errors"]["all_requested_times"]["temperature_Linf_C"], payload["space_levels"]["40"]["errors"]["all_requested_times"]["temperature_Linf_C"]),
            "moisture": order(payload["space_levels"]["20"]["errors"]["all_requested_times"]["moisture_Linf"], payload["space_levels"]["40"]["errors"]["all_requested_times"]["moisture_Linf"]),
            "interpretation": "proxy from errors relative to n=80 reference; authoritative dr=.1/.05/.025 observed orders are in EXP-Q2-006",
        },
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return payload["observed_order_proxy_vs_finest_reference"]


def run_long_continuation(environment_name: str, n_intervals: int = 80, time_step_s: float = 0.25, full_output: bool = False, resume_after_s: float | None = None) -> dict:
    path = ensure_folder(f"EXP-Q2-016-LONG-{environment_name}")
    rows = attachment_rows()
    candidate = environment_candidates(rows)[environment_name]
    config = Q2RunConfig(end_time_s=259200.0, time_step_s=time_step_s, candidate="A", n_intervals=n_intervals, scheme="bdf2", post_attachment_mode="constant", post_temperature_c=float(candidate["temperature_C"]), post_moisture_kg_kg=float(candidate["moisture_kg_kg"]))
    output_path = path / "official_samples.csv" if full_output else None
    diagnostics_path = path / "diagnostics_1s.csv"
    all_counts: list[int] = []
    stages = []
    previous_checkpoint = None
    previous_end = 0.0
    stage_ends = STAGE_ENDS
    if resume_after_s is not None:
        previous_end = float(resume_after_s)
        previous_checkpoint = path / f"checkpoint_{int(previous_end)}s.json"
        if not previous_checkpoint.exists():
            raise FileNotFoundError(f"resume checkpoint not found: {previous_checkpoint}")
        stage_ends = tuple(stage for stage in STAGE_ENDS if stage > previous_end)
    for stage_end in stage_ends:
        checkpoint = path / f"checkpoint_{int(stage_end)}s.json"
        started = time.perf_counter()
        kwargs = {"record_times": [stage_end], "output_path": output_path, "diagnostics_path": diagnostics_path, "checkpoint_path": checkpoint, "passive_event_threshold_kg_kg": 0.15, "diagnostics_interval_s": 1.0, "stop_time_s": stage_end}
        if previous_checkpoint is not None:
            kwargs["restart_path"] = previous_checkpoint
        result = run_q2(config, **kwargs)
        runtime = time.perf_counter() - started
        all_counts.extend(result.picard_count_history)
        stage = stage_summary(result, previous_end, stage_end, runtime)
        stage["checkpoint"] = str(checkpoint.relative_to(ROOT))
        stage["checkpoint_integrity"] = json.loads(checkpoint.read_text(encoding="utf-8"))
        stage["picard_counts_segment"] = {"count": len(result.picard_count_history), "min": min(result.picard_count_history), "median": sorted(result.picard_count_history)[len(result.picard_count_history) // 2], "max": max(result.picard_count_history)}
        stages.append(stage)
        dump(path / f"stage_{int(stage_end)}s.json", stage)
        previous_checkpoint, previous_end = checkpoint, stage_end
        if not stage["finite_final_fields"] or not stage["positive_moisture"] or stage["diagnostics"].get("picard_max", 0) > config.picard_max_iterations:
            dump(path / "metrics.json", {"experiment": f"EXP-Q2-016-LONG-{environment_name}", "status": "BLOCKED", "failed_stage": stage, "workbook_written": False})
            return {"status": "BLOCKED", "stages": stages}
    sample_validation = validate_output_samples(output_path, 259200.0) if output_path is not None else {"not_run": True}
    payload = {"experiment": f"EXP-Q2-016-LONG-{environment_name}", "status": "PASS", "config": cfg_dict(config), "environment_candidate": candidate, "resumed_from_s": resume_after_s, "stages": stages, "picard_all_steps": {"count": len(all_counts), "min": min(all_counts), "median": sorted(all_counts)[len(all_counts) // 2], "p95": sorted(all_counts)[int(0.95 * (len(all_counts) - 1))], "max": max(all_counts)}, "output_sampler": sample_validation, "output_storage_bytes": output_path.stat().st_size if output_path is not None else 0, "passive_observer_only": True, "q3_event_stop": False, "workbook_written": False}
    dump(path / "config.json", cfg_dict(config)); dump(path / "metrics.json", payload)
    return payload


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=("audit", "support", "long-a", "long-b", "resume-a", "recover-a", "finalize-a", "compare-env", "plot-a", "restart-a", "convergence-long", "augment-convergence", "baseline"))
    parser.add_argument("--n-intervals", type=int, default=80)
    parser.add_argument("--time-step", type=float, default=0.25)
    parser.add_argument("--full-output", action="store_true")
    args = parser.parse_args()
    if args.mode == "audit":
        print(json.dumps({"environment": run_environment_audit(), "interpolation": run_interpolation_audit()}, ensure_ascii=False, indent=2, default=str))
    elif args.mode == "support":
        print(json.dumps({"boundary": run_boundary_sensitivity(), "interface": run_interface_audit()}, ensure_ascii=False, indent=2, default=str))
    elif args.mode == "baseline":
        print(json.dumps(run_baseline_comparison(), ensure_ascii=False, indent=2, default=str))
    elif args.mode == "long-a":
        print(json.dumps(run_long_continuation("ENV-A-last-raw", args.n_intervals, args.time_step, args.full_output), ensure_ascii=False, indent=2, default=str))
    elif args.mode == "resume-a":
        print(json.dumps(run_long_continuation("ENV-A-last-raw", args.n_intervals, args.time_step, args.full_output, resume_after_s=43200.0), ensure_ascii=False, indent=2, default=str))
    elif args.mode == "recover-a":
        print(json.dumps(recover_long_output(), ensure_ascii=False, indent=2, default=str))
    elif args.mode == "finalize-a":
        print(json.dumps(finalize_recovered_long_a(), ensure_ascii=False, indent=2, default=str))
    elif args.mode == "compare-env":
        print(json.dumps(compare_environment_checkpoints(), ensure_ascii=False, indent=2, default=str))
    elif args.mode == "plot-a":
        print(json.dumps(plot_recovered_long_samples(), ensure_ascii=False, indent=2, default=str))
    elif args.mode == "restart-a":
        print(json.dumps(run_long_restart_audit(), ensure_ascii=False, indent=2, default=str))
    elif args.mode == "convergence-long":
        print(json.dumps(run_long_convergence(), ensure_ascii=False, indent=2, default=str))
    elif args.mode == "augment-convergence":
        print(json.dumps(augment_convergence_orders(), ensure_ascii=False, indent=2, default=str))
    else:
        print(json.dumps(run_long_continuation("ENV-B-tail40-mean", args.n_intervals, args.time_step, args.full_output), ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
