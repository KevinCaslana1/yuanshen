"""Remove non-lattice event rows from existing Q3/Q4 candidate workbooks.

This is a deliverable-only postprocess.  It reads the already generated
candidate workbooks, removes the exact threshold-event row from the official
60-second workbook lattice, and records hashes/row counts.  It never imports
or runs either Q3 or Q4 solver.
"""

from __future__ import annotations

import hashlib
import json
import math
import os
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RECORD = ROOT / "experiments/Q3_Q4_CANDIDATE/lattice_normalization.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def is_lattice_time(value: object) -> bool:
    if not isinstance(value, (int, float)) or not math.isfinite(float(value)):
        return False
    time_s = float(value)
    return time_s > 0.0 and abs(time_s / 60.0 - round(time_s / 60.0)) <= 1e-12


def normalize(path: Path, expected_last_s: float) -> dict[str, object]:
    before_sha = sha256(path)
    workbook = load_workbook(path, read_only=False, data_only=False)
    if len(workbook.worksheets) != 1:
        raise RuntimeError(f"expected one worksheet: {path}")
    sheet = workbook.active
    original_rows = sheet.max_row
    removed: list[dict[str, object]] = []
    for row_index in range(sheet.max_row, 1, -1):
        value = sheet.cell(row=row_index, column=1).value
        if not is_lattice_time(value):
            removed.append({"row": row_index, "time_s": value})
            sheet.delete_rows(row_index, 1)
    times = [sheet.cell(row=row_index, column=1).value for row_index in range(2, sheet.max_row + 1)]
    if not times or any(not is_lattice_time(value) for value in times):
        raise RuntimeError(f"non-lattice time remains: {path}")
    if any(abs(float(b) - float(a) - 60.0) > 1e-12 for a, b in zip(times, times[1:])):
        raise RuntimeError(f"time spacing is not 60 s: {path}")
    if abs(float(times[-1]) - expected_last_s) > 1e-12:
        raise RuntimeError(f"unexpected final lattice time: {path}")

    temporary = path.with_suffix(path.suffix + ".lattice.tmp")
    workbook.save(temporary)
    workbook.close()
    os.replace(temporary, path)
    after_sha = sha256(path)
    return {
        "path": str(path.relative_to(ROOT)).replace("\\", "/"),
        "before_sha256": before_sha,
        "after_sha256": after_sha,
        "original_rows": original_rows,
        "final_rows": len(times) + 1,
        "removed_rows": removed,
        "strict_60_s_lattice": True,
        "first_time_s": float(times[0]),
        "last_time_s": float(times[-1]),
        "expected_last_time_s": expected_last_s,
        "solver_rerun": False,
    }


def main() -> None:
    q3_summary = json.loads((ROOT / "experiments/Q3_PRODUCTION/q3_summary.json").read_text(encoding="utf-8"))
    q4_summary = json.loads((ROOT / "experiments/Q4_PRODUCTION/q4_summary.json").read_text(encoding="utf-8"))
    records = [
        normalize(ROOT / "deliverables/candidate/result3.xlsx", math.floor(q3_summary["t3_s"] / 60.0) * 60.0),
        normalize(ROOT / "deliverables/candidate/result4.xlsx", math.floor(q4_summary["t4_s"] / 60.0) * 60.0),
    ]
    payload = {
        "status": "PASS",
        "purpose": "official workbook 60 s lattice normalization; exact event rows remain paper/audit-only",
        "records": records,
        "solver_rerun": False,
    }
    RECORD.parent.mkdir(parents=True, exist_ok=True)
    RECORD.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
