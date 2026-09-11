"""Read-only attachment 1 input and configurable boundary interpolation."""

from __future__ import annotations

from bisect import bisect_right
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence, Tuple

from openpyxl import load_workbook

from src.common.paths import resolve_project_path


def celsius_to_kelvin(value_c: float) -> float:
    return float(value_c) + 273.15


def kelvin_to_celsius(value_k: float) -> float:
    return float(value_k) - 273.15


def cm_to_m(value_cm: float) -> float:
    return float(value_cm) / 100.0


@dataclass(frozen=True)
class BoundaryProvider:
    """Piecewise-linear provider that exactly reproduces every raw point."""

    times_s: Tuple[float, ...]
    temperatures_c: Tuple[float, ...]
    moistures_kg_kg: Tuple[float, ...]
    method: str = "linear"

    @classmethod
    def from_attachment1(cls, path: str | Path) -> "BoundaryProvider":
        resolved = resolve_project_path(path)
        workbook = load_workbook(resolved, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(min_row=2, values_only=True))
        workbook.close()
        if not rows:
            raise ValueError("attachment 1 contains no data rows")

        times = []
        temperatures = []
        moistures = []
        for row in rows:
            if len(row) < 3 or any(value is None for value in row[:3]):
                raise ValueError("attachment 1 contains an incomplete row")
            times.append(float(row[0]))
            temperatures.append(float(row[1]))
            moistures.append(float(row[2]))

        if any(right <= left for left, right in zip(times, times[1:])):
            raise ValueError("attachment 1 times must be strictly increasing")
        return cls(tuple(times), tuple(temperatures), tuple(moistures))

    @property
    def raw_count(self) -> int:
        return len(self.times_s)

    def _interpolate(self, values: Sequence[float], time_s: float) -> float:
        if self.method not in {"linear", "zero_order"}:
            raise ValueError(f"unsupported boundary interpolation method: {self.method}")
        if time_s < self.times_s[0] or time_s > self.times_s[-1]:
            raise ValueError(f"boundary time {time_s} s requires extrapolation")
        right = bisect_right(self.times_s, time_s)
        if right == 0:
            return float(values[0])
        if right == len(self.times_s):
            return float(values[-1])
        left = right - 1
        if self.times_s[left] == time_s:
            return float(values[left])
        if self.method == "zero_order":
            return float(values[left])
        fraction = (time_s - self.times_s[left]) / (self.times_s[right] - self.times_s[left])
        return float(values[left] + fraction * (values[right] - values[left]))

    def at(self, time_s: float) -> Tuple[float, float]:
        """Return `(temperature_c, moisture_kg_kg)` at a time in the raw range."""

        return (
            self._interpolate(self.temperatures_c, float(time_s)),
            self._interpolate(self.moistures_kg_kg, float(time_s)),
        )
