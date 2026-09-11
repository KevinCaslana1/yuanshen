"""Read-only Attachment 1 environment providers for Q2."""

from __future__ import annotations

from bisect import bisect_right
from dataclasses import dataclass, field
from hashlib import sha256
from pathlib import Path
from typing import Sequence, Tuple

from openpyxl import load_workbook

from src.common.paths import resolve_project_path


@dataclass(frozen=True)
class EnvironmentProvider:
    """Attachment 1 interpolation with explicit post-data behavior.

    ``linear`` is the reproducible default. ``pchip`` uses a small local
    monotone-cubic implementation, so the candidate comparison does not
    depend on an optional scientific-Python package. Both methods reproduce
    every raw Attachment 1 knot exactly.
    """

    times_s: Tuple[float, ...]
    temperatures_c: Tuple[float, ...]
    moistures_kg_kg: Tuple[float, ...]
    method: str = "linear"
    post_attachment_mode: str = "raise"
    post_temperature_c: float = 50.0
    post_moisture_kg_kg: float = 0.05
    source_path: str = ""
    source_sha256: str = ""
    _temperature_slopes: Tuple[float, ...] = field(default=(), init=False, repr=False, compare=False)
    _moisture_slopes: Tuple[float, ...] = field(default=(), init=False, repr=False, compare=False)

    @classmethod
    def from_attachment1(
        cls,
        path: str | Path,
        *,
        method: str = "linear",
        post_attachment_mode: str = "raise",
        post_temperature_c: float = 50.0,
        post_moisture_kg_kg: float = 0.05,
    ) -> "EnvironmentProvider":
        resolved = resolve_project_path(path)
        raw = resolved.read_bytes()
        workbook = load_workbook(resolved, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(min_row=2, values_only=True))
        workbook.close()
        if not rows:
            raise ValueError("attachment 1 contains no data rows")
        times, temperatures, moistures = [], [], []
        for row in rows:
            if len(row) < 3 or any(value is None for value in row[:3]):
                raise ValueError("attachment 1 contains an incomplete row")
            times.append(float(row[0]))
            temperatures.append(float(row[1]))
            moistures.append(float(row[2]))
        if any(right <= left for left, right in zip(times, times[1:])):
            raise ValueError("attachment 1 times must be strictly increasing")
        return cls(
            tuple(times), tuple(temperatures), tuple(moistures), method,
            post_attachment_mode, float(post_temperature_c),
            float(post_moisture_kg_kg), str(resolved), sha256(raw).hexdigest(),
        )

    def __post_init__(self) -> None:
        if self.method not in {"linear", "pchip"}:
            raise ValueError("method must be linear or pchip")
        if self.post_attachment_mode not in {"raise", "constant"}:
            raise ValueError("post_attachment_mode must be raise or constant")
        if len(self.times_s) < 2 or len(self.times_s) != len(self.temperatures_c) or len(self.times_s) != len(self.moistures_kg_kg):
            raise ValueError("environment arrays have inconsistent lengths")
        if any(right <= left for left, right in zip(self.times_s, self.times_s[1:])):
            raise ValueError("environment times must be strictly increasing")
        if self.method == "pchip":
            # Build once here to fail closed on malformed data and to keep the
            # interpolation path deterministic during long runs.
            object.__setattr__(self, "_temperature_slopes", self._pchip_slopes(self.temperatures_c))
            object.__setattr__(self, "_moisture_slopes", self._pchip_slopes(self.moistures_kg_kg))

    @property
    def raw_count(self) -> int:
        return len(self.times_s)

    @property
    def last_time_s(self) -> float:
        return self.times_s[-1]

    def _linear(self, values: Sequence[float], time_s: float) -> float:
        right = bisect_right(self.times_s, time_s)
        if right == 0:
            return float(values[0])
        if right == len(self.times_s):
            return float(values[-1])
        left = right - 1
        if self.times_s[left] == time_s:
            return float(values[left])
        fraction = (time_s - self.times_s[left]) / (self.times_s[right] - self.times_s[left])
        return float(values[left] + fraction * (values[right] - values[left]))

    def _pchip(self, values: Sequence[float], time_s: float, slopes: Sequence[float]) -> float:
        if time_s == self.times_s[0]:
            return float(values[0])
        if time_s == self.times_s[-1]:
            return float(values[-1])
        right = bisect_right(self.times_s, time_s)
        left = right - 1
        h = self.times_s[right] - self.times_s[left]
        x = (time_s - self.times_s[left]) / h
        y0, y1 = values[left], values[right]
        h00 = (1.0 + 2.0 * x) * (1.0 - x) ** 2
        h10 = x * (1.0 - x) ** 2
        h01 = x * x * (3.0 - 2.0 * x)
        h11 = x * x * (x - 1.0)
        return float(h00 * y0 + h10 * h * slopes[left] + h01 * y1 + h11 * h * slopes[right])

    def _pchip_slopes(self, values: Sequence[float]) -> Tuple[float, ...]:
        """Return Fritsch-Carlson/PCHIP knot slopes without hidden smoothing."""
        n = len(self.times_s)
        if len(values) != n:
            raise ValueError("PCHIP values do not match environment times")
        h = [self.times_s[i + 1] - self.times_s[i] for i in range(n - 1)]
        delta = [(values[i + 1] - values[i]) / h[i] for i in range(n - 1)]
        if n == 2:
            return (delta[0], delta[0])

        def endpoint(h0, h1, d0, d1):
            value = ((2.0 * h0 + h1) * d0 - h0 * d1) / (h0 + h1)
            if value * d0 <= 0.0:
                return 0.0
            if d0 * d1 < 0.0 and abs(value) > 3.0 * abs(d0):
                return 3.0 * d0
            return value

        slopes = [0.0] * n
        slopes[0] = endpoint(h[0], h[1], delta[0], delta[1])
        slopes[-1] = endpoint(h[-1], h[-2], delta[-1], delta[-2])
        for i in range(1, n - 1):
            if delta[i - 1] * delta[i] <= 0.0:
                slopes[i] = 0.0
            else:
                w1 = 2.0 * h[i] + h[i - 1]
                w2 = h[i] + 2.0 * h[i - 1]
                slopes[i] = (w1 + w2) / (w1 / delta[i - 1] + w2 / delta[i])
        return tuple(slopes)

    def at(self, time_s: float) -> Tuple[float, float]:
        time_s = float(time_s)
        if time_s < self.times_s[0]:
            raise ValueError(f"boundary time {time_s} s precedes Attachment 1")
        if time_s > self.times_s[-1]:
            if self.post_attachment_mode == "raise":
                raise ValueError(f"boundary time {time_s} s requires an unapproved post-attachment rule")
            return self.post_temperature_c, self.post_moisture_kg_kg
        if self.method == "linear":
            return self._linear(self.temperatures_c, time_s), self._linear(self.moistures_kg_kg, time_s)
        return self._pchip(self.temperatures_c, time_s, self._temperature_slopes), self._pchip(self.moistures_kg_kg, time_s, self._moisture_slopes)
