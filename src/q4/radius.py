"""Attachment 2 prescribed-radius reader and PCHIP mapping."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from functools import cached_property
from pathlib import Path
from typing import Tuple

from openpyxl import load_workbook

from src.q4.pchip import PchipInterpolator


@dataclass(frozen=True)
class RadiusLaw:
    times_s: Tuple[float, ...]
    radii_cm: Tuple[float, ...]
    source_path: str
    source_sha256: str

    @classmethod
    def from_attachment2(cls, path: str | Path) -> "RadiusLaw":
        resolved = Path(path).resolve()
        raw = resolved.read_bytes()
        wb = load_workbook(resolved, read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(min_row=2, values_only=True))
        wb.close()
        times = tuple(float(row[0]) for row in rows if row[0] is not None and row[1] is not None)
        radii = tuple(float(row[1]) for row in rows if row[0] is not None and row[1] is not None)
        if len(times) < 2 or any(b <= a for a, b in zip(times, times[1:])):
            raise ValueError("Attachment 2 time nodes must increase")
        if any(b > a + 1e-12 for a, b in zip(radii, radii[1:])):
            raise ValueError("Attachment 2 radius must be nonincreasing")
        return cls(times, radii, str(resolved), hashlib.sha256(raw).hexdigest())

    @property
    def last_time_s(self) -> float:
        return self.times_s[-1]

    @property
    def last_radius_cm(self) -> float:
        return self.radii_cm[-1]

    @cached_property
    def _pchip(self) -> PchipInterpolator:
        return PchipInterpolator(self.times_s, self.radii_cm)

    def radius_cm(self, time_s: float) -> float:
        time_s = float(time_s)
        if time_s >= self.last_time_s:
            return self.last_radius_cm
        if time_s < self.times_s[0]:
            raise ValueError("time precedes Attachment 2")
        return self._pchip(time_s)

    def radius_rate_cm_s(self, time_s: float) -> float:
        if time_s >= self.last_time_s:
            return 0.0
        if time_s < self.times_s[0]:
            raise ValueError("time precedes Attachment 2")
        return self._pchip.derivative(time_s)
