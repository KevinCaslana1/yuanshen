"""Regression tests for the Q3/Q4 candidate package."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from src.q4.properties import conductivity, cp, density, diffusivity
from src.q4.radius import RadiusLaw


ROOT = Path(__file__).resolve().parents[1]


def test_q3_candidate_threshold_and_endpoint_contract() -> None:
    summary = json.loads((ROOT / "experiments/Q3_PRODUCTION/q3_summary.json").read_text(encoding="utf-8"))
    assert summary["prior_not_satisfied"] is True
    assert summary["endpoint_satisfied"] is True
    assert summary["coarse_bracket_s"] == [206935.0, 206936.0]
    assert summary["endpoint_critical_radius_cm"] == 0.0


def test_q3_result3_shape_and_monotonic_time() -> None:
    wb = load_workbook(ROOT / "deliverables/candidate/result3.xlsx", read_only=True, data_only=False)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    assert len(rows) == 3449
    assert len(rows[0]) == 22
    assert list(rows[0][1:]) == [i / 10 for i in range(21)]
    t3 = json.loads((ROOT / "experiments/Q3_PRODUCTION/q3_summary.json").read_text(encoding="utf-8"))["t3_s"]
    assert rows[-1][0] == int(t3 // 60) * 60
    assert all(rows[i][0] > 0 and rows[i][0] % 60 == 0 for i in range(1, len(rows)))
    assert all(rows[i][0] < rows[i + 1][0] for i in range(1, len(rows) - 1))


def test_q4_result4_uses_strict_60_second_lattice() -> None:
    wb = load_workbook(ROOT / "deliverables/candidate/result4.xlsx", read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    assert len(rows) == 3185
    assert len(rows[0]) == 22
    t4 = json.loads((ROOT / "experiments/Q4_PRODUCTION/q4_summary.json").read_text(encoding="utf-8"))["t4_s"]
    assert rows[-1][0] == int(t4 // 60) * 60
    assert all(rows[i][0] > 0 and rows[i][0] % 60 == 0 for i in range(1, len(rows)))
    assert all(rows[i][0] < rows[i + 1][0] for i in range(1, len(rows) - 1))


def test_q4_radius_nodes_are_exact_and_monotone() -> None:
    law = RadiusLaw.from_attachment2(ROOT / "A题/附件/附件2.xlsx")
    assert max(abs(law.radius_cm(t) - r) for t, r in zip(law.times_s, law.radii_cm)) == 0.0
    samples = [law.radius_cm(i * law.last_time_s / 100.0) for i in range(101)]
    assert all(b <= a + 1e-12 for a, b in zip(samples, samples[1:]))
    assert law.radius_cm(law.last_time_s + 10.0) == law.last_radius_cm


def test_q4_appendix4_properties_are_finite_and_positive() -> None:
    c = 0.15
    t_k = 323.15
    assert density(c) > 0
    assert cp(c) > 0
    assert conductivity(c) > 0
    assert diffusivity(c, t_k) > 0
